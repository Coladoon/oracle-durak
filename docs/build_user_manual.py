from __future__ import annotations

import math
import os
import textwrap
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
ASSETS = DOCS / "guide_assets"
OUTPUT = DOCS / "Руководство_по_запуску_и_игре_Дурак_Oracle.docx"

BLUE = "1F4E78"
LIGHT_BLUE = "D9EAF7"
PALE_BLUE = "EDF5FB"
GREEN = "2F6B4F"
LIGHT_GREEN = "E2F0D9"
ORANGE = "C65911"
LIGHT_ORANGE = "FCE4D6"
RED = "A61B1B"
LIGHT_RED = "F4CCCC"
GRAY = "666666"
LIGHT_GRAY = "E7E6E6"
VERY_LIGHT_GRAY = "F5F5F5"
BLACK = "000000"
WHITE = "FFFFFF"


def get_font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    candidates = [
        "/System/Library/Fonts/Supplemental/Arial Bold.ttf" if bold else "/System/Library/Fonts/Supplemental/Arial.ttf",
        "/System/Library/Fonts/Supplemental/Arial Unicode.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf" if bold else "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    ]
    for candidate in candidates:
        if os.path.exists(candidate):
            return ImageFont.truetype(candidate, size)
    return ImageFont.load_default()


FONT_46 = get_font(46, True)
FONT_34 = get_font(34, True)
FONT_30 = get_font(30, True)
FONT_26 = get_font(26, False)
FONT_24 = get_font(24, True)
FONT_22 = get_font(22, False)
FONT_20 = get_font(20, False)
FONT_18 = get_font(18, False)


def rgb(hex_color: str) -> tuple[int, int, int]:
    return tuple(int(hex_color[i : i + 2], 16) for i in (0, 2, 4))


def rounded_box(draw: ImageDraw.ImageDraw, xy, fill, outline=BLUE, width=3, radius=22):
    draw.rounded_rectangle(xy, radius=radius, fill=rgb(fill), outline=rgb(outline), width=width)


def wrapped_lines(text: str, font, max_width: int, draw: ImageDraw.ImageDraw) -> list[str]:
    lines: list[str] = []
    for source_line in text.split("\n"):
        words = source_line.split()
        if not words:
            lines.append("")
            continue
        current = words[0]
        for word in words[1:]:
            trial = current + " " + word
            if draw.textbbox((0, 0), trial, font=font)[2] <= max_width:
                current = trial
            else:
                lines.append(current)
                current = word
        lines.append(current)
    return lines


def draw_centered_text(draw, box, text, font, fill=BLACK, spacing=8):
    x1, y1, x2, y2 = box
    lines = wrapped_lines(text, font, x2 - x1 - 32, draw)
    heights = [draw.textbbox((0, 0), line or " ", font=font)[3] for line in lines]
    total = sum(heights) + spacing * (len(lines) - 1)
    y = y1 + (y2 - y1 - total) / 2
    for line, height in zip(lines, heights):
        bbox = draw.textbbox((0, 0), line, font=font)
        x = x1 + (x2 - x1 - (bbox[2] - bbox[0])) / 2
        draw.text((x, y), line, font=font, fill=rgb(fill))
        y += height + spacing


def arrow(draw, start, end, color=BLUE, width=8, head=22):
    draw.line([start, end], fill=rgb(color), width=width)
    angle = math.atan2(end[1] - start[1], end[0] - start[0])
    left = (
        end[0] - head * math.cos(angle - math.pi / 6),
        end[1] - head * math.sin(angle - math.pi / 6),
    )
    right = (
        end[0] - head * math.cos(angle + math.pi / 6),
        end[1] - head * math.sin(angle + math.pi / 6),
    )
    draw.polygon([end, left, right], fill=rgb(color))


def canvas(title: str) -> tuple[Image.Image, ImageDraw.ImageDraw]:
    image = Image.new("RGB", (1600, 900), "white")
    draw = ImageDraw.Draw(image)
    draw.rectangle((0, 0, 1600, 92), fill=rgb(BLUE))
    draw.text((55, 20), title, font=FONT_46, fill="white")
    return image, draw


def build_diagrams() -> None:
    ASSETS.mkdir(parents=True, exist_ok=True)

    image, draw = canvas("Кто что делает в Oracle")
    boxes = [
        ((65, 180, 500, 700), LIGHT_RED, RED, "SYSTEM / SYS", "Администрирует базу\n\n• создаёт DURAK_P1…DURAK_P6\n• задаёт и сбрасывает пароли\n• разблокирует учётные записи\n• проверяет DBA_USERS"),
        ((580, 180, 1015, 700), LIGHT_BLUE, BLUE, "DURAK_OWNER", "Владеет проектом\n\n• запускает install.sql\n• выполняет reinstall.sql\n• выдаёт права игрокам\n• запускает tests/run_all.sql\n• проверяет объекты и планировщик"),
        ((1095, 180, 1530, 700), LIGHT_GREEN, GREEN, "DURAK_P1…P6", "Играют\n\n• регистрируются\n• создают или входят в партию\n• атакуют и защищаются\n• читают защищённые представления\n• не видят чужие карты"),
    ]
    for box, fill, outline, heading, body in boxes:
        rounded_box(draw, box, fill, outline, 4)
        draw_centered_text(draw, (box[0] + 20, box[1] + 28, box[2] - 20, box[1] + 110), heading, FONT_34, outline)
        draw_centered_text(draw, (box[0] + 30, box[1] + 125, box[2] - 30, box[3] - 30), body, FONT_22, BLACK, 11)
    draw_centered_text(draw, (120, 750, 1480, 850), "Главное правило: смотрите имя активного подключения в правом верхнем углу Worksheet, а не только название вкладки.", FONT_26, RED)
    image.save(ASSETS / "01_roles.png")

    image, draw = canvas("Маршрут от пустой базы до первой партии")
    steps = [
        ("1", "Запустить\nDocker Oracle"),
        ("2", "Подключить\nDURAK_OWNER"),
        ("3", "Выполнить\ninstall.sql"),
        ("4", "Создать 6\nпользователей"),
        ("5", "Выдать им\nправа"),
        ("6", "Создать 6\nподключений"),
        ("7", "Создать, войти\nи запустить игру"),
        ("8", "Играть и\nсмотреть журнал"),
    ]
    positions = [(55 + (i % 4) * 385, 165 + (i // 4) * 330) for i in range(8)]
    for i, ((number, label), (x, y)) in enumerate(zip(steps, positions)):
        box = (x, y, x + 330, y + 205)
        rounded_box(draw, box, PALE_BLUE if i < 4 else LIGHT_GREEN, BLUE if i < 4 else GREEN, 4)
        draw.ellipse((x + 18, y + 20, x + 88, y + 90), fill=rgb(BLUE))
        draw_centered_text(draw, (x + 18, y + 20, x + 88, y + 90), number, FONT_34, WHITE)
        draw_centered_text(draw, (x + 95, y + 20, x + 315, y + 185), label, FONT_24, BLACK)
        if i not in (3, 7):
            arrow(draw, (x + 330, y + 102), (x + 370, y + 102), BLUE, 6, 18)
    arrow(draw, (1510, 370), (1510, 470), BLUE, 6, 18)
    draw_centered_text(draw, (120, 785, 1480, 865), "Если шаг не завершён успешно, к следующему не переходите.", FONT_26, ORANGE)
    image.save(ASSETS / "02_launch_flow.png")

    image, draw = canvas("Шесть игроков: ход идёт по кругу")
    cx, cy = 800, 500
    draw.ellipse((610, 310, 990, 690), fill=rgb(PALE_BLUE), outline=rgb(BLUE), width=5)
    draw_centered_text(draw, (640, 350, 960, 650), "ИГРА\n\nстол\nкозырь\nталон\nжурнал", FONT_30, BLUE)
    seats = [
        (800, 145, "1. Михаил\nАТАКУЕТ", LIGHT_ORANGE, ORANGE),
        (1260, 290, "2. Анна\nЗАЩИЩАЕТСЯ", LIGHT_RED, RED),
        (1260, 670, "3. Борис\nожидает", LIGHT_GREEN, GREEN),
        (800, 790, "4. Вера\nожидает", PALE_BLUE, BLUE),
        (340, 670, "5. Глеб\nожидает", PALE_BLUE, BLUE),
        (340, 290, "6. Даша\nожидает", PALE_BLUE, BLUE),
    ]
    for x, y, label, fill, outline in seats:
        box = (x - 175, y - 72, x + 175, y + 72)
        rounded_box(draw, box, fill, outline, 4, 28)
        draw_centered_text(draw, box, label, FONT_22, BLACK)
    # Clockwise direction markers.
    for start, end in [((1010, 175), (1140, 235)), ((1400, 420), (1400, 545)), ((1100, 765), (970, 815)), ((630, 815), (500, 765)), ((200, 545), (200, 420)), ((460, 235), (590, 175))]:
        arrow(draw, start, end, GRAY, 5, 16)
    draw_centered_text(draw, (60, 105, 500, 175), "По часовой стрелке", FONT_24, GRAY)
    image.save(ASSETS / "03_six_players.png")

    image, draw = canvas("Один понятный цикл ручного хода")
    items = [
        ((65, 210, 385, 560), "1", "ОБНОВИТЬ", "SELECT …\nFROM V_GAME_CONSOLE", LIGHT_BLUE, BLUE),
        ((455, 210, 775, 560), "2", "ПОНЯТЬ", "Кто ходит?\nКакая фаза?\nЧто разрешено?", LIGHT_ORANGE, ORANGE),
        ((845, 210, 1165, 560), "3", "ВЫПОЛНИТЬ", "Одну команду\nDURAK_CONSOLE.PLAY", LIGHT_GREEN, GREEN),
        ((1235, 210, 1535, 560), "4", "ПРОВЕРИТЬ", "PRINT message\nи снова обновить\nконсоль", PALE_BLUE, BLUE),
    ]
    for box, number, heading, body, fill, outline in items:
        rounded_box(draw, box, fill, outline, 4)
        draw.ellipse((box[0] + 20, box[1] + 25, box[0] + 92, box[1] + 97), fill=rgb(outline))
        draw_centered_text(draw, (box[0] + 20, box[1] + 25, box[0] + 92, box[1] + 97), number, FONT_34, WHITE)
        draw_centered_text(draw, (box[0] + 105, box[1] + 22, box[2] - 15, box[1] + 110), heading, FONT_30, outline)
        draw_centered_text(draw, (box[0] + 25, box[1] + 125, box[2] - 25, box[3] - 25), body, FONT_24, BLACK)
    for x in (405, 795, 1185):
        arrow(draw, (x, 385), (x + 35, 385), BLUE, 7, 18)
    arrow(draw, (1380, 590), (1380, 735), GRAY, 6, 18)
    arrow(draw, (1380, 735), (220, 735), GRAY, 6, 18)
    arrow(draw, (220, 735), (220, 590), GRAY, 6, 18)
    draw_centered_text(draw, (400, 660, 1200, 800), "Повторять после каждого успешного действия", FONT_30, GRAY)
    image.save(ASSETS / "04_game_loop.png")

    image, draw = canvas("Фазы партии и допустимые действия")
    states = [
        ("LOBBY", "игроки входят"),
        ("WAIT_ATTACK", "ХОД <карта>"),
        ("WAIT_DEFENSE", "БИТО / ВЗЯТЬ / ПЕРЕВОД"),
        ("WAIT_THROW", "ПОДКИНУТЬ / ПАС"),
        ("ROUND_END", "добор и смена ролей"),
        ("FINISHED", "определён дурак"),
    ]
    xs = [40, 300, 570, 870, 1160, 1325]
    widths = [220, 230, 260, 250, 255, 235]
    for i, ((state, label), x, w) in enumerate(zip(states, xs, widths)):
        y = 250 if i < 5 else 610
        box = (x, y, x + w, y + 220)
        fill = LIGHT_RED if state == "FINISHED" else PALE_BLUE
        outline = RED if state == "FINISHED" else BLUE
        rounded_box(draw, box, fill, outline, 4)
        draw_centered_text(draw, (x + 10, y + 22, x + w - 10, y + 95), state, FONT_24, outline)
        draw_centered_text(draw, (x + 18, y + 98, x + w - 18, y + 198), label, FONT_20, BLACK)
        if i < 4:
            arrow(draw, (x + w, y + 110), (xs[i + 1] - 12, y + 110), BLUE, 5, 15)
    arrow(draw, (1285, 470), (1420, 585), BLUE, 5, 15)
    arrow(draw, (1185, 490), (470, 555), GRAY, 5, 15)
    draw.text((690, 525), "если игра продолжается", font=FONT_20, fill=rgb(GRAY))
    draw_centered_text(draw, (100, 720, 1150, 845), "Точную текущую фазу всегда берите из TURN_TEXT или поля PHASE, а не угадывайте по предыдущему ходу.", FONT_26, ORANGE)
    image.save(ASSETS / "05_states.png")


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color="B7B7B7", size="5") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=90, start=100, bottom=90, end=100):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def add_page_number(paragraph):
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    header = table.rows[0]
    set_repeat_table_header(header)
    for i, value in enumerate(headers):
        cell = header.cells[i]
        cell.text = value
        set_cell_shading(cell, BLUE)
        set_cell_border(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for run in paragraph.runs:
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.size = Pt(9)
    for row_values in rows:
        row = table.add_row()
        prevent_row_split(row)
        for i, value in enumerate(row_values):
            cell = row.cells[i]
            cell.text = str(value)
            set_cell_border(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if len(table.rows) % 2 == 1:
                set_cell_shading(cell, "F8FAFC")
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.size = Pt(8.8)
    if widths:
        for row in table.rows:
            for cell, width in zip(row.cells, widths):
                cell.width = Inches(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_code(doc: Document, code: str, title: str | None = None):
    if title:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(title)
        run.bold = True
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor.from_string(BLUE)
    p = doc.add_paragraph()
    p.style = doc.styles["Code Block"]
    p.paragraph_format.keep_together = True
    for idx, line in enumerate(code.strip("\n").splitlines()):
        if idx:
            p.add_run().add_break()
        p.add_run(line)
    return p


def add_note(doc: Document, label: str, text: str, color: str = BLUE):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.35)
    p.paragraph_format.right_indent = Cm(0.2)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(7)
    r1 = p.add_run(label + " ")
    r1.bold = True
    r1.font.color.rgb = RGBColor.from_string(color)
    p.add_run(text)
    p_pr = p._p.get_or_add_pPr()
    border = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), color)
    border.append(left)
    p_pr.append(border)


def add_figure(doc: Document, filename: str, caption: str, width=6.75):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    shape = p.add_run().add_picture(str(ASSETS / filename), width=Inches(width))
    shape._inline.docPr.set("descr", caption)
    shape._inline.docPr.set("title", caption)
    cap = doc.add_paragraph(caption)
    cap.style = doc.styles["Caption"]
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(8)


def bullet(doc: Document, text: str, level=0):
    style = "List Bullet" if level == 0 else "List Bullet 2"
    p = doc.add_paragraph(text, style=style)
    p.paragraph_format.space_after = Pt(2)
    return p


NUMBER_COUNTER = 0


def number(doc: Document, text: str):
    global NUMBER_COUNTER
    NUMBER_COUNTER += 1
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.55)
    p.paragraph_format.first_line_indent = Cm(-0.55)
    p.paragraph_format.space_after = Pt(3)
    prefix = p.add_run(f"{NUMBER_COUNTER}. ")
    prefix.bold = True
    p.add_run(text)
    return p


def heading(doc: Document, text: str, level=1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True
    return p


def new_section(doc: Document, title: str, intro: str | None = None):
    global NUMBER_COUNTER
    NUMBER_COUNTER = 0
    doc.add_page_break()
    heading(doc, title, 1)
    if intro:
        p = doc.add_paragraph(intro)
        p.style = doc.styles["Lead"]


def build_document() -> None:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.67)
    section.bottom_margin = Inches(0.62)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    normal.font.size = Pt(10.2)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08

    title_style = styles["Title"]
    title_style.font.name = "Arial"
    title_style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    title_style.font.size = Pt(30)
    title_style.font.bold = True
    title_style.font.color.rgb = RGBColor(0, 0, 0)
    title_style.paragraph_format.space_after = Pt(12)

    for name, size, before, after in (("Heading 1", 20, 3, 9), ("Heading 2", 14, 9, 5), ("Heading 3", 11.5, 7, 3)):
        style = styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    if "Lead" not in styles:
        lead = styles.add_style("Lead", WD_STYLE_TYPE.PARAGRAPH)
    else:
        lead = styles["Lead"]
    lead.font.name = "Arial"
    lead.font.size = Pt(11.5)
    lead.font.color.rgb = RGBColor.from_string(GRAY)
    lead.paragraph_format.space_after = Pt(10)

    if "Code Block" not in styles:
        code_style = styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
    else:
        code_style = styles["Code Block"]
    code_style.font.name = "Courier New"
    code_style._element.rPr.rFonts.set(qn("w:eastAsia"), "Courier New")
    code_style.font.size = Pt(8.2)
    code_style.paragraph_format.left_indent = Cm(0.5)
    code_style.paragraph_format.right_indent = Cm(0.2)
    code_style.paragraph_format.space_before = Pt(3)
    code_style.paragraph_format.space_after = Pt(7)
    code_style.paragraph_format.line_spacing = 1.0

    caption = styles["Caption"]
    caption.font.name = "Arial"
    caption.font.size = Pt(8.5)
    caption.font.italic = True
    caption.font.color.rgb = RGBColor.from_string(GRAY)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run("Дурак на Oracle • руководство пользователя • ")
    footer_run.font.size = Pt(8)
    footer_run.font.color.rgb = RGBColor.from_string(GRAY)
    add_page_number(footer)

    # Cover.
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(52)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("РУКОВОДСТВО ПО ЗАПУСКУ И ИГРЕ")
    run.bold = True
    run.font.size = Pt(13)
    run.font.color.rgb = RGBColor.from_string(BLUE)
    p = doc.add_paragraph("«Дурак» только на Oracle", style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph("Docker Oracle Free • SQL Developer • 2–6 игроков")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].font.size = Pt(15)
    p.runs[0].font.color.rgb = RGBColor.from_string(GRAY)
    p.paragraph_format.space_after = Pt(24)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cover_shape = p.add_run().add_picture(str(ASSETS / "02_launch_flow.png"), width=Inches(6.7))
    cover_shape._inline.docPr.set("descr", "Восемь этапов запуска проекта от Docker до ручной игры")
    cover_shape._inline.docPr.set("title", "Маршрут запуска проекта")
    p = doc.add_paragraph("Версия для защиты минимального и хорошего уровней требований")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].bold = True
    p.runs[0].font.size = Pt(11)
    p.paragraph_format.space_before = Pt(8)
    p = doc.add_paragraph("Подготовлено для Михаила • 25 сентября 2026")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].font.size = Pt(9.5)
    p.runs[0].font.color.rgb = RGBColor.from_string(GRAY)

    doc.add_page_break()
    heading(doc, "Как пользоваться этим руководством", 1)
    doc.add_paragraph(
        "Руководство ведёт от запуска пустой базы до ручной партии на шесть человек. "
        "Каждый шаг отвечает на четыре вопроса: какое подключение открыть, что выполнить, "
        "какой результат считать правильным и что делать при ошибке."
    )
    add_note(doc, "Главная проверка.", "Перед любой командой смотрите имя активного подключения в правом верхнем углу Worksheet и выполняйте SELECT USER FROM dual. Ошибки чаще всего возникают не из-за SQL, а из-за запуска в другом подключении.", RED)
    heading(doc, "Краткий маршрут", 2)
    add_table(
        doc,
        ["Этап", "Подключение", "Результат"],
        [
            ["1. Oracle", "Terminal", "Контейнер запущен, база READY"],
            ["2. Проект", "DURAK_OWNER", "Объекты установлены и VALID"],
            ["3. Пользователи", "SYSTEM", "DURAK_P1…DURAK_P6 существуют и открыты"],
            ["4. Права", "DURAK_OWNER", "Игроки могут выполнять API и читать views"],
            ["5. Подключения", "Каждый DURAK_Pn", "Test: Success для всех шести"],
            ["6. Партия", "P1 создаёт, P2–P6 входят", "Статус ACTIVE, карты розданы"],
            ["7. Игра", "Текущий игрок", "Команды проходят, журнал растёт"],
            ["8. Проверка", "DURAK_OWNER", "78 успешно, 0 ошибок"],
        ],
        [1.25, 1.8, 3.65],
    )
    heading(doc, "Обозначения", 2)
    bullet(doc, "<GAME_ID> — номер вашей партии, например 2102. Угловые скобки вводить не надо.")
    bullet(doc, "<пароль> — ваш пароль. Не копируйте это слово буквально.")
    bullet(doc, "F5 — Run Script. Используйте его для файлов, VARIABLE, PRINT, @ и блоков с символом /.")
    bullet(doc, "Ctrl+Enter — запуск одного выделенного SQL-оператора; для установочных сценариев не подходит.")
    bullet(doc, "Все пути в примерах относятся к текущему проекту в каталоге «Игра Дурак только на Oracle».")

    new_section(doc, "1. Структура игры", "Проект работает полностью внутри Oracle: SQL Developer служит только терминалом ввода и просмотра.")
    add_figure(doc, "01_roles.png", "Рисунок 1 — роли подключений и границы ответственности")
    heading(doc, "1.1. Что происходит после команды", 2)
    add_table(
        doc,
        ["Уровень", "Объект", "Назначение"],
        [
            ["Ввод", "DURAK_CONSOLE", "Принимает короткие команды на русском языке."],
            ["Доступ", "DURAK_API", "Определяет игрока по SESSION_USER, проверяет участие и возвращает сообщение."],
            ["Игровой автомат", "DURAK_ENGINE", "Блокирует партию, проверяет фазу и атомарно меняет состояние."],
            ["Правила", "DURAK_RULES", "Проверяет бой карт, козырь и лимиты стола."],
            ["Раздача", "DURAK_RANDOM", "Строит воспроизводимый порядок колоды по seed."],
            ["Таймеры", "DURAK_MAINTENANCE", "Выполняет автоматические действия и завершает простой."],
            ["Чтение", "V_*", "Маскирует чужие руки и показывает разрешённую информацию."],
        ],
        [1.15, 1.65, 3.9],
    )
    heading(doc, "1.2. Основные каталоги", 2)
    add_table(
        doc,
        ["Путь", "Что находится"],
        [
            ["install.sql", "Чистая установка таблиц, пакетов, представлений и таймера."],
            ["reinstall.sql", "Удаление старых объектов и повторная установка; игровые данные удаляются."],
            ["sql/", "Исходный Oracle SQL и PL/SQL."],
            ["admin/", "Выдача игрокам минимальных прав."],
            ["demo/", "Готовые сценарии создания, подключения, старта и игры."],
            ["tests/", "78 автоматических проверок и тест конкурентности."],
            ["docs/", "Инструкции и матрица выполнения требований."],
            ["reports/", "Отчёты ЛР1 и ЛР2."],
            ["versions/", "Архив старых редакций; для основной демонстрации не используется."],
        ],
        [1.55, 5.15],
    )

    new_section(doc, "2. Запуск Oracle в Docker", "Этот раздел нужен при первом запуске или после перезагрузки компьютера.")
    heading(doc, "2.1. Первый запуск контейнера", 2)
    doc.add_paragraph("Откройте Terminal и выполните команду, заменив оба пароля своими значениями:")
    add_code(doc, """docker run -d \\
  --name durak-oracle \\
  -p 1521:1521 \\
  -e ORACLE_PASSWORD='<пароль_SYS>' \\
  -e APP_USER=DURAK_OWNER \\
  -e APP_USER_PASSWORD='<пароль_DURAK_OWNER>' \\
  -v durak-oracle-data:/opt/oracle/oradata \\
  gvenzl/oracle-free:latest""")
    add_note(doc, "Безопасность.", "Не добавляйте реальные пароли в Git, отчёт или скриншоты. В SQL Developer пароль можно сохранить только на своём компьютере.", ORANGE)
    heading(doc, "2.2. Дождаться готовности", 2)
    add_code(doc, "docker logs -f durak-oracle")
    doc.add_paragraph("Продолжайте только после строки:")
    add_code(doc, "DATABASE IS READY TO USE!")
    heading(doc, "2.3. Последующие запуски", 2)
    add_code(doc, "docker start durak-oracle\ndocker ps")
    add_table(
        doc,
        ["Ситуация", "Проверка или действие"],
        [
            ["Контейнер не виден в docker ps", "Выполнить docker start durak-oracle."],
            ["Имя уже занято", "Контейнер уже создан; не выполнять docker run повторно."],
            ["Порт 1521 занят", "Остановить другой Oracle или изменить проброс порта и то же значение указать в SQL Developer."],
            ["SQL Developer подключается слишком рано", "Посмотреть logs и дождаться READY."],
        ],
        [2.35, 4.35],
    )

    new_section(doc, "3. Подключения в SQL Developer", "Все подключения используют одну базу FREEPDB1, но имеют разные полномочия.")
    heading(doc, "3.1. Общие сетевые поля", 2)
    add_table(
        doc,
        ["Поле", "Значение", "Важно"],
        [
            ["Connection Type", "Basic", "Не TNS и не Cloud Wallet."],
            ["Hostname", "localhost", "Контейнер опубликован на локальном компьютере."],
            ["Port", "1521", "Или ваш порт слева в -p <порт>:1521."],
            ["Address type", "Service name", "Выберите радиокнопку Service name."],
            ["Service name", "FREEPDB1", "Не SID xe и не служба FREE."],
            ["Role", "default", "SYS — единственное исключение: SYSDBA."],
        ],
        [1.55, 1.8, 3.35],
    )
    heading(doc, "3.2. Три группы подключений", 2)
    add_table(
        doc,
        ["Имя подключения", "Username", "Для чего"],
        [
            ["system-freepdb1", "SYSTEM", "Создание, разблокировка и сброс паролей игроков."],
            ["durak-owner", "DURAK_OWNER", "Установка проекта, выдача прав, тесты."],
            ["durak-p1…durak-p6", "DURAK_P1…DURAK_P6", "Ручная игра от имени разных участников."],
        ],
        [1.7, 1.85, 3.15],
    )
    add_note(doc, "Критически важно.", "Название вкладки может вводить в заблуждение. Реальное подключение показано в выпадающем списке справа вверху редактора Worksheet.", RED)
    heading(doc, "3.3. Проверка активного подключения", 2)
    doc.add_paragraph("В каждом новом Worksheet сначала выполните:")
    add_code(doc, "SELECT USER AS current_user,\n       SYS_CONTEXT('USERENV', 'CON_NAME') AS container_name\nFROM dual;")
    add_table(
        doc,
        ["Окно", "Ожидаемый CURRENT_USER", "Ожидаемый CONTAINER_NAME"],
        [
            ["Администратор", "SYSTEM", "FREEPDB1"],
            ["Владелец", "DURAK_OWNER", "FREEPDB1"],
            ["Игрок 1", "DURAK_P1", "FREEPDB1"],
        ],
        [1.7, 2.35, 2.65],
    )

    new_section(doc, "4. Установка проекта", "Установку выполняет только DURAK_OWNER.")
    number(doc, "Откройте подключение durak-owner.")
    number(doc, "Выберите File → Open и откройте install.sql из корня проекта.")
    number(doc, "Проверьте справа вверху: активное подключение DURAK_OWNER.")
    number(doc, "Нажмите F5 (Run Script) и дождитесь завершения.")
    add_note(doc, "Почему F5.", "Файл содержит @, VARIABLE, PRINT, SHOW ERRORS и разделители /. Ctrl+Enter пытается выполнить только один оператор и ломает сценарий.", ORANGE)
    heading(doc, "4.1. Проверка установки", 2)
    add_code(doc, """SELECT object_name, object_type, status
FROM user_objects
WHERE object_type IN ('PACKAGE', 'PACKAGE BODY', 'VIEW')
  AND status <> 'VALID'
ORDER BY object_type, object_name;""")
    doc.add_paragraph("Правильный результат: no rows selected.")
    add_code(doc, """SELECT job_name, enabled, state
FROM user_scheduler_jobs
WHERE job_name = 'DURAK_MAINTENANCE_JOB';""")
    doc.add_paragraph("Правильный результат: одна строка, ENABLED = TRUE.")
    heading(doc, "4.2. Повторная установка", 2)
    add_code(doc, "@reinstall.sql")
    add_note(doc, "Осторожно.", "reinstall.sql удаляет текущие партии, игроков проекта и историю. Используйте только для чистой демонстрации или восстановления схемы.", RED)

    new_section(doc, "5. Создание шести Oracle-пользователей", "Это административный этап: выполнять в подключении SYSTEM, не в DURAK_OWNER.")
    heading(doc, "5.1. Открыть правильное окно", 2)
    number(doc, "В панели Connections щёлкните правой кнопкой system-freepdb1.")
    number(doc, "Выберите Open SQL Worksheet.")
    number(doc, "Убедитесь, что справа вверху написано system-freepdb1.")
    number(doc, "Выполните SELECT USER и проверьте SYSTEM / FREEPDB1.")
    heading(doc, "5.2. Создать пользователей", 2)
    doc.add_paragraph("Ниже показаны примерные пароли для локальной учебной базы. При необходимости замените их своими:")
    add_code(doc, """CREATE USER DURAK_P1 IDENTIFIED BY "Player1_2026!";
GRANT CREATE SESSION TO DURAK_P1;
CREATE USER DURAK_P2 IDENTIFIED BY "Player2_2026!";
GRANT CREATE SESSION TO DURAK_P2;
CREATE USER DURAK_P3 IDENTIFIED BY "Player3_2026!";
GRANT CREATE SESSION TO DURAK_P3;
CREATE USER DURAK_P4 IDENTIFIED BY "Player4_2026!";
GRANT CREATE SESSION TO DURAK_P4;
CREATE USER DURAK_P5 IDENTIFIED BY "Player5_2026!";
GRANT CREATE SESSION TO DURAK_P5;
CREATE USER DURAK_P6 IDENTIFIED BY "Player6_2026!";
GRANT CREATE SESSION TO DURAK_P6;""")
    heading(doc, "5.3. Если пользователи уже существуют", 2)
    add_code(doc, """ALTER USER DURAK_P1 IDENTIFIED BY "Player1_2026!" ACCOUNT UNLOCK;
GRANT CREATE SESSION TO DURAK_P1;""")
    doc.add_paragraph("Повторите ALTER USER для P2…P6 с соответствующими паролями.")
    heading(doc, "5.4. Проверить результат", 2)
    add_code(doc, """SELECT username, account_status
FROM dba_users
WHERE username IN ('DURAK_P1','DURAK_P2','DURAK_P3',
                   'DURAK_P4','DURAK_P5','DURAK_P6')
ORDER BY username;""")
    doc.add_paragraph("Ожидается шесть строк со статусом OPEN.")
    add_note(doc, "Если ORA-00942 на DBA_USERS.", "Вы не в SYSTEM/SYS. Создайте новый Worksheet именно из system-freepdb1 и повторите SELECT USER.", RED)
    add_note(doc, "Если ORA-01031 на CREATE USER.", "Команда запущена от DURAK_OWNER или игрока. Перейдите в SYSTEM.", RED)

    new_section(doc, "6. Выдача игровых прав", "Теперь вернитесь в DURAK_OWNER. SYSTEM на этом шаге не нужен.")
    number(doc, "Откройте новый Worksheet от durak-owner.")
    number(doc, "Проверьте SELECT USER: должен быть DURAK_OWNER.")
    number(doc, "Вставьте команды ниже и нажмите F5.")
    add_code(doc, """@"/Users/desertik/Documents/ChatGPT/Игра Дурак только на Oracle/admin/grant_player_access.sql" DURAK_P1
@"/Users/desertik/Documents/ChatGPT/Игра Дурак только на Oracle/admin/grant_player_access.sql" DURAK_P2
@"/Users/desertik/Documents/ChatGPT/Игра Дурак только на Oracle/admin/grant_player_access.sql" DURAK_P3
@"/Users/desertik/Documents/ChatGPT/Игра Дурак только на Oracle/admin/grant_player_access.sql" DURAK_P4
@"/Users/desertik/Documents/ChatGPT/Игра Дурак только на Oracle/admin/grant_player_access.sql" DURAK_P5
@"/Users/desertik/Documents/ChatGPT/Игра Дурак только на Oracle/admin/grant_player_access.sql" DURAK_P6""")
    doc.add_paragraph("В Script Output должны появиться сообщения Grant succeeded. Сценарий выдаёт только нужные права на DURAK_API, DURAK_CONSOLE и защищённые представления.")
    add_note(doc, "Не запускать Ctrl+Enter.", "Команда @ является командой сценария SQL*Plus; запускайте весь блок клавишей F5.", ORANGE)

    new_section(doc, "7. Создание шести подключений игроков", "Для каждого DURAK_Pn создаётся отдельное подключение с одинаковыми сетевыми параметрами.")
    add_table(
        doc,
        ["Name", "Username", "Password (пример)", "Service"],
        [
            ["durak-p1", "DURAK_P1", "Player1_2026!", "FREEPDB1"],
            ["durak-p2", "DURAK_P2", "Player2_2026!", "FREEPDB1"],
            ["durak-p3", "DURAK_P3", "Player3_2026!", "FREEPDB1"],
            ["durak-p4", "DURAK_P4", "Player4_2026!", "FREEPDB1"],
            ["durak-p5", "DURAK_P5", "Player5_2026!", "FREEPDB1"],
            ["durak-p6", "DURAK_P6", "Player6_2026!", "FREEPDB1"],
        ],
        [1.45, 1.5, 2.1, 1.65],
    )
    doc.add_paragraph("Для каждого подключения нажмите Test. Только после Status: Success нажмите Save и Connect.")
    heading(doc, "Если Test возвращает ORA-01017", 2)
    bullet(doc, "Проверьте, что Username и пароль относятся к одному игроку.")
    bullet(doc, "Проверьте Service name = FREEPDB1; не SID xe.")
    bullet(doc, "В SYSTEM выполните ALTER USER … IDENTIFIED BY … ACCOUNT UNLOCK.")
    bullet(doc, "Введите новый пароль вручную — SQL Developer мог сохранить старый.")
    add_code(doc, "ALTER USER DURAK_P1 IDENTIFIED BY " + '"Player1_2026!"' + " ACCOUNT UNLOCK;")

    new_section(doc, "8. Подготовка шести игровых окон", "Откройте по одному Worksheet для P1…P6. Для удобства расположите вкладки рядом или подпишите их цветами.")
    doc.add_paragraph("В каждом из шести окон один раз выполните этот блок клавишей F5:")
    add_code(doc, """SET SERVEROUTPUT ON SIZE UNLIMITED
ALTER SESSION SET CURRENT_SCHEMA = DURAK_OWNER;
VARIABLE message VARCHAR2(1000)""")
    doc.add_paragraph("Затем проверяйте, что SESSION_USER остался игроком, хотя CURRENT_SCHEMA стал владельцем:")
    add_code(doc, """SELECT SYS_CONTEXT('USERENV','SESSION_USER') AS session_user,
       SYS_CONTEXT('USERENV','CURRENT_SCHEMA') AS current_schema
FROM dual;""")
    add_table(
        doc,
        ["Окно", "SESSION_USER", "CURRENT_SCHEMA"],
        [
            ["P1", "DURAK_P1", "DURAK_OWNER"],
            ["P2", "DURAK_P2", "DURAK_OWNER"],
            ["…", "…", "…"],
            ["P6", "DURAK_P6", "DURAK_OWNER"],
        ],
        [1.2, 2.35, 2.75],
    )
    add_note(doc, "Почему это безопасно.", "CURRENT_SCHEMA меняет только поиск имён объектов. Игрок всё равно определяется по SESSION_USER и не получает прав владельца.", GREEN)

    new_section(doc, "9. Создание партии на шесть игроков", "Для наглядной шестиместной демонстрации рекомендуется колода 52 карты: после раздачи остаётся 16 карт в талоне.")
    add_figure(doc, "03_six_players.png", "Рисунок 2 — места игроков и направление очереди")
    heading(doc, "9.1. Игрок P1 регистрируется и создаёт партию", 2)
    doc.add_paragraph("В окне DURAK_P1 выполните:")
    add_code(doc, """VARIABLE player_id NUMBER
VARIABLE game_id NUMBER

BEGIN
  durak_api.register_player(
    p_player_id    => :player_id,
    p_message      => :message,
    p_display_name => 'Михаил'
  );
END;
/

BEGIN
  durak_api.create_game(
    p_game_id          => :game_id,
    p_message          => :message,
    p_deck_size        => 52,
    p_game_variant     => 'PODKIDNOY',
    p_first_move_mode  => 'LOWEST_TRUMP',
    p_turn_timeout_sec => 0,
    p_idle_timeout_min => 120,
    p_seed_text        => 'six-player-test-001'
  );
END;
/

PRINT player_id
PRINT game_id
PRINT message""")
    doc.add_paragraph("Запишите GAME_ID. Далее в примерах используется 2102, но вы подставляете свой номер.")
    heading(doc, "9.2. Игроки P2…P6 регистрируются и входят", 2)
    doc.add_paragraph("В каждом окне изменяйте только имя. Пример для P2:")
    add_code(doc, """VARIABLE player_id NUMBER
VARIABLE seat_no NUMBER

BEGIN
  durak_api.register_player(
    p_player_id    => :player_id,
    p_message      => :message,
    p_display_name => 'Анна'
  );
END;
/

BEGIN
  durak_api.join_game(
    p_game_id => 2102,
    p_seat_no => :seat_no,
    p_message => :message
  );
END;
/

PRINT player_id
PRINT seat_no
PRINT message""")
    add_table(
        doc,
        ["Окно", "Имя для регистрации", "GAME_ID"],
        [
            ["DURAK_P2", "Анна", "тот же"],
            ["DURAK_P3", "Борис", "тот же"],
            ["DURAK_P4", "Вера", "тот же"],
            ["DURAK_P5", "Глеб", "тот же"],
            ["DURAK_P6", "Даша", "тот же"],
        ],
        [1.8, 2.45, 2.05],
    )
    heading(doc, "9.3. Проверить состав до старта", 2)
    add_code(doc, "SELECT game_id, seat_no, display_name, player_status\nFROM v_game_status\nWHERE game_id = 2102\nORDER BY seat_no;")
    doc.add_paragraph("Ожидается шесть игроков. Если меньше — не запускайте игру, подключите недостающих.")
    heading(doc, "9.4. Создатель запускает раздачу", 2)
    doc.add_paragraph("Только в окне P1:")
    add_code(doc, """BEGIN
  durak_api.start_game(
    p_game_id => 2102,
    p_message => :message
  );
END;
/
PRINT message""")

    new_section(doc, "10. Как читать экран игры", "После старта главным экраном становится представление V_GAME_CONSOLE.")
    add_code(doc, """SELECT game_text,
       turn_text,
       hand_text,
       table_text,
       players_text,
       command_hint,
       last_event_text
FROM v_game_console
WHERE game_id = 2102;""")
    add_table(
        doc,
        ["Поле", "Как читать"],
        [
            ["GAME_TEXT", "Номер, режим, колода, статус, козырь и остаток талона."],
            ["TURN_TEXT", "Фаза, атакующий, защитник и конкретный текущий исполнитель."],
            ["HAND_TEXT", "Только ваша открытая рука; в другом подключении содержимое будет другим."],
            ["TABLE_TEXT", "Пары атака/защита с номерами 1, 2, 3…"],
            ["PLAYERS_TEXT", "Места, статусы и количество карт у всех участников."],
            ["COMMAND_HINT", "Разрешённая именно вам команда на текущем шаге."],
            ["LAST_EVENT_TEXT", "Последнее событие, помогающее понять, что только что изменилось."],
        ],
        [1.6, 5.1],
    )
    add_note(doc, "Если no rows selected.", "Вы указали неправильный GAME_ID, игрок не является участником либо партия завершена/удалена. Сначала проверьте номер и V_GAME_STATUS.", ORANGE)
    heading(doc, "Подробные представления", 2)
    add_code(doc, """SELECT * FROM v_game_status WHERE game_id = 2102 ORDER BY seat_no;
SELECT * FROM v_hand_mine WHERE game_id = 2102 ORDER BY rank_value, suit_code;
SELECT * FROM v_hand_public WHERE game_id = 2102 ORDER BY seat_no;
SELECT * FROM v_table WHERE game_id = 2102 ORDER BY pair_no;""")

    new_section(doc, "11. Как выполнять ходы", "Начиная с этого раздела, после каждого действия повторяйте один и тот же цикл.")
    add_figure(doc, "04_game_loop.png", "Рисунок 3 — цикл одного ручного игрового действия")
    add_figure(doc, "05_states.png", "Рисунок 4 — основные фазы игрового автомата")
    doc.add_page_break()
    heading(doc, "11.1. Алгоритм после каждого хода", 2)
    number(doc, "Во всех окнах обновите V_GAME_CONSOLE.")
    number(doc, "По TURN_TEXT найдите текущего игрока.")
    number(doc, "Перейдите именно на вкладку его подключения.")
    number(doc, "Выберите карту из HAND_TEXT и команду из COMMAND_HINT.")
    number(doc, "Выполните ровно одну команду и PRINT message.")
    number(doc, "Снова обновите V_GAME_CONSOLE и прочитайте LAST_EVENT_TEXT.")
    heading(doc, "11.2. Команды", 2)
    add_table(
        doc,
        ["Команда", "Когда", "Пример"],
        [
            ["ХОД <карта>", "Первая атака раунда", "ХОД 8H"],
            ["БИТО <карта> <пара>", "Защитник покрывает атаку", "БИТО 9H 1"],
            ["ПОДКИНУТЬ <карта>", "Разрешено атакующей стороне", "ПОДКИНУТЬ 8C"],
            ["ПЕРЕВОД <карта>", "Переводной режим, до первой защиты", "ПЕРЕВОД 8D"],
            ["ВЗЯТЬ", "Защитник не может или не хочет отбиваться", "ВЗЯТЬ"],
            ["ПАС", "Подбрасывающий завершает участие в подбрасывании", "ПАС"],
        ],
        [2.0, 3.15, 1.55],
    )
    add_code(doc, """EXEC durak_console.play('ХОД 8H', :message)
PRINT message

EXEC durak_console.play('БИТО 9H 1', :message)
PRINT message""")
    doc.add_paragraph("Можно отправить только код карты — консоль сама выберет действие по фазе:")
    add_code(doc, "EXEC durak_console.play('8H', :message)\nPRINT message")
    add_note(doc, "Для защиты лучше писать явно.", "При нескольких парах используйте БИТО <карта> <номер пары>, чтобы не закрыть не ту атаку.", ORANGE)
    heading(doc, "11.3. Коды карт", 2)
    add_table(
        doc,
        ["Элемент", "Коды", "Примеры"],
        [
            ["Ранги 36", "6, 7, 8, 9, 10, J, Q, K, A", "6C, 10H, QS, AD"],
            ["Дополнительно 52", "2, 3, 4, 5", "2D, 5S"],
            ["Масти", "C трефы; D бубны; H червы; S пики", "QC — дама треф"],
        ],
        [1.45, 3.45, 1.8],
    )

    new_section(doc, "12. Правила атаки, защиты и подбрасывания", "Движок проверяет правила автоматически, но игроку полезно понимать причину отказа.")
    heading(doc, "12.1. Чем бить карту", 2)
    add_table(
        doc,
        ["Атака", "Защита", "Результат"],
        [
            ["8H", "9H", "Можно: старше той же масти."],
            ["KH", "6S, если S козырь", "Можно: козырь бьёт некозырь."],
            ["8S, если S козырь", "9S", "Можно: старший козырь."],
            ["8H", "9D", "Нельзя: другая некозырная масть."],
            ["9H", "8H", "Нельзя: карта младше."],
            ["9S, если S козырь", "AH", "Нельзя: некозырь не бьёт козырь."],
        ],
        [1.85, 2.25, 2.6],
    )
    heading(doc, "12.2. Подбрасывание", 2)
    bullet(doc, "Подбрасывать можно только ранг, уже присутствующий на столе среди атак или защит.")
    bullet(doc, "Количество пар не превышает MAX_PAIRS: обычно 6 для колоды 36 и 8 для 52.")
    bullet(doc, "Лимит также не превышает число карт защитника на начало атаки.")
    bullet(doc, "В игре на 3–6 человек право подбрасывать передаётся по очереди, пропуская вышедших.")
    heading(doc, "12.3. Завершение раунда", 2)
    bullet(doc, "Все атаки побиты и атакующая сторона пасует — стол уходит в отбой.")
    bullet(doc, "Защитник выбирает ВЗЯТЬ — все карты стола переходят в его руку.")
    bullet(doc, "После раунда сначала добирают атакующие по кругу, затем защитник.")
    bullet(doc, "При пустом талоне игрок без карт выходит; последний с картами становится дураком.")

    new_section(doc, "13. Параметры и режимы партии", "Основная версия поддерживает все требования минимального и хорошего уровней.")
    add_table(
        doc,
        ["Параметр", "Значения", "Рекомендация для показа"],
        [
            ["P_DECK_SIZE", "36 или 52", "36 для 2–4; 52 для 6 игроков."],
            ["P_GAME_VARIANT", "PODKIDNOY / PEREVODNOY", "Начать с PODKIDNOY."],
            ["P_FIRST_MOVE_MODE", "LOWEST_TRUMP / SEEDED_RANDOM", "LOWEST_TRUMP показывает классическое правило."],
            ["P_MAX_PAIRS", "по умолчанию 6 / 8", "Оставить NULL/по умолчанию."],
            ["P_TURN_TIMEOUT_SEC", "0 или секунды", "0 при ручной демонстрации."],
            ["P_IDLE_TIMEOUT_MIN", "минуты простоя", "120 при длинной проверке."],
            ["P_SEED_TEXT", "произвольная строка", "Записать seed для повторения раздачи."],
        ],
        [2.0, 2.45, 2.25],
    )
    heading(doc, "13.1. Переводная партия", 2)
    add_code(doc, """BEGIN
  durak_api.create_game(
    p_game_id          => :game_id,
    p_message          => :message,
    p_deck_size        => 36,
    p_game_variant     => 'PEREVODNOY',
    p_first_move_mode  => 'SEEDED_RANDOM',
    p_turn_timeout_sec => 0,
    p_idle_timeout_min => 120,
    p_seed_text        => 'transfer-demo-001'
  );
END;
/""")
    bullet(doc, "Перевод разрешён только до первой защиты.")
    bullet(doc, "Защитник добавляет карту того же достоинства, что первая атака.")
    bullet(doc, "Следующий активный игрок становится новым защитником.")
    bullet(doc, "Если лимит стола или рука следующего защитника не позволяют перевод, команда отклоняется.")
    heading(doc, "13.2. Детерминированный seed", 2)
    doc.add_paragraph("Одинаковые параметры и одинаковый seed дают одинаковый порядок колоды. Для сравнения создайте две чистые партии с тем же seed и одинаковым составом/местами игроков.")

    new_section(doc, "14. Таймеры и автоматические действия", "Для ручной игры таймер лучше отключить; для проверки требования включите отдельную партию.")
    add_table(
        doc,
        ["Фаза", "Действие при истечении"],
        [
            ["Ожидание первой атаки", "Движок выбирает допустимую карту."],
            ["Ожидание защиты", "Защитник автоматически берёт."],
            ["Ожидание подбрасывания", "Текущий подбрасывающий пасует."],
            ["Длительный простой партии", "Партия получает статус EXPIRED с причиной."],
        ],
        [2.55, 4.15],
    )
    heading(doc, "14.1. Проверка фонового задания", 2)
    add_code(doc, """SELECT job_name, enabled, state, last_start_date
FROM user_scheduler_jobs
WHERE job_name = 'DURAK_MAINTENANCE_JOB';""")
    heading(doc, "14.2. Безопасный тест", 2)
    doc.add_paragraph("Создайте отдельную игру с p_turn_timeout_sec => 10. После старта не выполняйте действие 10–20 секунд, обновите V_GAME_CONSOLE и V_LOG_TEXT. В журнале должна появиться причина TIMEOUT.")

    new_section(doc, "15. Журнал, реплей, Daily и табло", "Эти функции демонстрируют требования хорошего уровня.")
    heading(doc, "15.1. Читаемый журнал", 2)
    add_code(doc, "SELECT log_line\nFROM v_log_text\nWHERE game_id = 2102\nORDER BY event_no;")
    doc.add_paragraph("Пример смысла строк: кто атаковал, чем отбились, кто взял, когда закончился раунд и сколько карт добрали.")
    heading(doc, "15.2. Пошаговый реплей", 2)
    add_code(doc, """SELECT replay_step_no,
       event_no,
       event_type,
       actor_name,
       event_card_code,
       moved_card_code,
       from_zone,
       to_zone,
       phase_before,
       phase_after,
       event_message
FROM v_replay
WHERE game_id = 2102
ORDER BY replay_step_no;""")
    add_note(doc, "Маскирование.", "Во время активной партии чужие закрытые события скрыты. После завершения участники видят полный реплей.", GREEN)
    heading(doc, "15.3. Daily", 2)
    doc.add_paragraph("В окне игрока-создателя запустите готовый сценарий через F5:")
    add_code(doc, "@demo/05_daily.sql DURAK_OWNER Михаил 2026-09-25")
    add_code(doc, """SELECT * FROM v_daily ORDER BY daily_date DESC, daily_id DESC;
SELECT * FROM v_daily_results
WHERE daily_id = <DAILY_ID>
ORDER BY daily_place, finish_place;""")
    heading(doc, "15.4. Табло", 2)
    add_code(doc, """SELECT display_name,
       games_played,
       wins,
       losses,
       draws,
       avg_duration_sec,
       avg_rounds
FROM v_leaderboard
ORDER BY wins DESC, losses, display_name;""")

    new_section(doc, "16. Автоматическое тестирование", "Полный прогон подтверждает требования минимального и хорошего уровней.")
    number(doc, "Откройте новый Worksheet от DURAK_OWNER.")
    number(doc, "Убедитесь, что активных ручных партий, мешающих тестам, нет.")
    number(doc, "Запустите блок ниже клавишей F5.")
    add_code(doc, "SET SERVEROUTPUT ON SIZE UNLIMITED\n@tests/run_all.sql")
    doc.add_paragraph("Ожидаемый итог:")
    add_code(doc, "--- Итого: успешно 78, ошибок 0 ---")
    add_table(
        doc,
        ["Группа", "Что подтверждается"],
        [
            ["Карты", "52 карты, короткая колода 36, уникальные коды."],
            ["Правила", "Масть, козырь, лимиты, недопустимые параметры."],
            ["Seed", "Повторяемость раздачи и Daily seed."],
            ["Доступ", "Своя рука открыта, чужая скрыта, журнал маскируется."],
            ["Игровой автомат", "Полный раунд, добор, взятие, отбой, простои."],
            ["Режимы", "2–6 игроков, 36/52, перевод, выход игрока."],
        ],
        [2.0, 4.7],
    )
    heading(doc, "16.1. Конкурентность двух сеансов", 2)
    doc.add_paragraph("Подробный сценарий находится в tests/concurrency/README.md. Ожидается: первая сессия удерживает блокировку, а вторая получает контролируемый отказ ORA-20702 вместо двойного хода.")

    new_section(doc, "17. Диагностика ошибок", "Не переустанавливайте проект сразу. Сначала определите, на каком этапе возникла ошибка.")
    add_table(
        doc,
        ["Сообщение", "Причина", "Что сделать"],
        [
            ["ORA-01017", "Неверный пользователь/пароль или учётная запись заблокирована.", "SYSTEM: ALTER USER … IDENTIFIED BY … ACCOUNT UNLOCK; затем ввести пароль заново."],
            ["ORA-01031", "Команда требует административных прав.", "CREATE/ALTER USER выполнять как SYSTEM; установку — как DURAK_OWNER."],
            ["ORA-00942 на DBA_USERS", "Worksheet фактически не SYSTEM.", "Проверить список подключения справа вверху и SELECT USER."],
            ["ORA-01403", "Ожидалась строка данных, но подготовка/справочник не установлены.", "Проверить install.sql и наличие 52 строк в DURAK_CARD."],
            ["ORA-04063 / ORA-06508", "Тело пакета невалидно.", "Показать USER_ERRORS, исправить/переустановить, затем проверить VALID."],
            ["no rows selected", "Фильтр не нашёл строку; иногда это ожидаемо.", "Для проверки invalid — хорошо; для игры проверить GAME_ID и участие."],
            ["У игрока нет карты 10C", "Карта отсутствует в вашей руке.", "Взять код из HAND_TEXT или V_HAND_MINE."],
            ["Сейчас нельзя начинать атаку", "Не ваша очередь или другая фаза.", "Прочитать TURN_TEXT и COMMAND_HINT, перейти к текущему игроку."],
            ["Игрок уже участвует…", "У SESSION_USER уже есть активная партия.", "Завершить/отменить прежнюю партию или использовать другого Oracle-пользователя."],
            ["ORA-20702", "Партию сейчас изменяет другой сеанс.", "Обновить состояние и повторить после завершения чужой операции."],
        ],
        [1.75, 2.45, 2.5],
    )
    heading(doc, "17.1. Команды быстрой проверки", 2)
    add_code(doc, """SELECT USER, SYS_CONTEXT('USERENV','CON_NAME') FROM dual;

SELECT COUNT(*) AS card_count FROM durak_card;

SELECT object_name, object_type, status
FROM user_objects
WHERE status <> 'VALID';

SELECT name, type, line, position, text
FROM user_errors
ORDER BY name, sequence;""")

    new_section(doc, "18. Сценарий демонстрации преподавателю", "Эта последовательность показывает все обязательные функции без лишних расширений уровня «отлично».")
    steps = [
        "Показать Docker-контейнер и подключение FREEPDB1.",
        "В DURAK_OWNER показать отсутствие INVALID-объектов и включённый scheduler job.",
        "Открыть два разных подключения и показать, что V_HAND_MINE возвращает разные руки.",
        "Создать подкидную игру 36 карт на двух игроков с известным seed.",
        "Показать младший козырь, первого атакующего, талон и по шесть карт.",
        "Выполнить атаку, правильную защиту, подбрасывание и пас или взятие.",
        "Показать добор до шести и уменьшение талона.",
        "Показать V_LOG_TEXT и V_REPLAY.",
        "Создать шестиместную игру 52 карты и показать шесть участников.",
        "Создать отдельную переводную игру и выполнить перевод.",
        "Создать игру с коротким таймером и показать автоматическое действие в протоколе.",
        "Показать Daily и V_DAILY_RESULTS.",
        "Показать V_LEADERBOARD после завершённой партии.",
        "Запустить tests/run_all.sql и показать 78/78.",
    ]
    for step in steps:
        number(doc, step)
    heading(doc, "Контрольный чек-лист перед защитой", 2)
    checklist_rows = [
        ["□", "Контейнер запускается и сообщает READY."],
        ["□", "DURAK_OWNER подключается к FREEPDB1."],
        ["□", "Все объекты VALID, scheduler включён."],
        ["□", "DURAK_P1…P6 имеют статус OPEN и Test: Success."],
        ["□", "Права выданы из DURAK_OWNER."],
        ["□", "Игра создаётся, участники входят, старт проходит."],
        ["□", "V_GAME_CONSOLE показывает руку, стол и подсказку."],
        ["□", "Ручная команда изменяет состояние и попадает в журнал."],
        ["□", "Daily, replay и leaderboard доступны."],
        ["□", "Автотесты завершаются 78 успешно, 0 ошибок."],
    ]
    add_table(doc, ["", "Проверка"], checklist_rows, [0.45, 6.25])

    new_section(doc, "Приложение A. Шпаргалка игрока", "Эту страницу удобно держать открытой во время ручной партии.")
    heading(doc, "1. В начале окна", 2)
    add_code(doc, """SET SERVEROUTPUT ON SIZE UNLIMITED
ALTER SESSION SET CURRENT_SCHEMA = DURAK_OWNER;
VARIABLE message VARCHAR2(1000)""")
    heading(doc, "2. Посмотреть состояние", 2)
    add_code(doc, """SELECT game_text, turn_text, hand_text, table_text,
       players_text, command_hint, last_event_text
FROM v_game_console
WHERE game_id = 2102;""")
    heading(doc, "3. Выполнить одну команду", 2)
    add_code(doc, """EXEC durak_console.play('ХОД 8H', :message)
EXEC durak_console.play('БИТО 9H 1', :message)
EXEC durak_console.play('ПОДКИНУТЬ 8C', :message)
EXEC durak_console.play('ПЕРЕВОД 8D', :message)
EXEC durak_console.play('ВЗЯТЬ', :message)
EXEC durak_console.play('ПАС', :message)
PRINT message""")
    heading(doc, "4. Посмотреть историю", 2)
    add_code(doc, "SELECT log_line FROM v_log_text\nWHERE game_id = 2102 ORDER BY event_no;")
    add_note(doc, "Если непонятно, что делать.", "Не угадывайте: обновите V_GAME_CONSOLE, прочитайте TURN_TEXT и COMMAND_HINT, затем перейдите в окно указанного игрока.", BLUE)

    doc.core_properties.title = "Руководство по запуску и игре «Дурак» только на Oracle"
    doc.core_properties.subject = "Oracle SQL Developer, Docker, ручная игра 2–6 игроков"
    doc.core_properties.author = "Михаил"
    doc.core_properties.keywords = "Oracle, SQL, PL/SQL, Дурак, SQL Developer, Docker"
    doc.core_properties.comments = "Подробная инструкция по проекту минимального и хорошего уровней."

    DOCS.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_diagrams()
    build_document()

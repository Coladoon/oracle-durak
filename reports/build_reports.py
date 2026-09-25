from __future__ import annotations

from pathlib import Path
from typing import Iterable, Sequence

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Cm, Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
REPORTS = ROOT / "reports"
ASSETS = REPORTS / "assets"
TEMPLATE_LR1 = Path("/Users/desertik/Downloads/Артамонова_Козлов_ЛР1.docx")
TEMPLATE_LR2 = Path("/Users/desertik/Downloads/Артамонова_Козлов_ЛР2 .docx")
OUT_LR1 = REPORTS / "Михаил_ЛР1_Дурак_Oracle.docx"
OUT_LR2 = REPORTS / "Михаил_ЛР2_Дурак_Oracle.docx"

FONT_REGULAR = Path("/System/Library/Fonts/Supplemental/Times New Roman.ttf")
FONT_BOLD = Path("/System/Library/Fonts/Supplemental/Times New Roman Bold.ttf")


def font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    return ImageFont.truetype(str(FONT_BOLD if bold else FONT_REGULAR), size)


def draw_arrow(draw: ImageDraw.ImageDraw, start: tuple[int, int], end: tuple[int, int], color="#40566f", width=4):
    draw.line([start, end], fill=color, width=width)
    x2, y2 = end
    x1, y1 = start
    import math

    angle = math.atan2(y2 - y1, x2 - x1)
    length = 16
    spread = 0.55
    p1 = (x2 - length * math.cos(angle - spread), y2 - length * math.sin(angle - spread))
    p2 = (x2 - length * math.cos(angle + spread), y2 - length * math.sin(angle + spread))
    draw.polygon([end, p1, p2], fill=color)


def draw_box(draw: ImageDraw.ImageDraw, xy, title: str, lines: Sequence[str], fill: str, outline: str = "#36495d"):
    x1, y1, x2, y2 = xy
    draw.rounded_rectangle(xy, radius=14, fill=fill, outline=outline, width=3)
    draw.rounded_rectangle((x1, y1, x2, y1 + 48), radius=14, fill=outline, outline=outline, width=2)
    draw.rectangle((x1, y1 + 30, x2, y1 + 50), fill=outline)
    draw.text(((x1 + x2) / 2, y1 + 24), title, font=font(25, True), fill="white", anchor="mm")
    y = y1 + 65
    for line in lines:
        draw.text((x1 + 16, y), line, font=font(19), fill="#1f2b38")
        y += 27


def create_diagrams() -> None:
    ASSETS.mkdir(parents=True, exist_ok=True)

    # Концептуальная схема данных для первой лабораторной работы.
    img = Image.new("RGB", (1800, 1050), "white")
    d = ImageDraw.Draw(img)
    d.text((900, 42), "Концептуальная модель данных игры «Дурак»", font=font(38, True), fill="#172b4d", anchor="mm")
    boxes = {
        "PLAYER": (60, 150, 370, 320, "#f9e2e8", ["player_id", "Oracle-учётная запись", "локаль, часовой пояс"]),
        "GAME": (510, 130, 870, 350, "#dcecff", ["game_id, seed", "вариант, колода, лимиты", "фаза, атакующий, защитник"]),
        "GAME_PLAYER": (60, 430, 410, 650, "#e5f5e0", ["место за столом", "роль и статус", "число карт, место выхода"]),
        "GAME_CARD": (510, 470, 870, 690, "#fff1cc", ["карта партии", "зона: рука/талон/стол", "позиция и владелец"]),
        "ROUND": (1010, 130, 1325, 325, "#eee5ff", ["номер раунда", "атакующий / защитник", "результат раунда"]),
        "TABLE_PAIR": (1435, 130, 1750, 325, "#e6f3ff", ["карта атаки", "карта защиты", "номер пары"]),
        "EVENT": (1010, 470, 1325, 690, "#fde7d9", ["полный протокол", "инициатор и причина", "снимок состояния"]),
        "DAILY": (1435, 470, 1750, 665, "#e4f7f4", ["дата и seed", "версия правил", "результаты дня"]),
        "ERROR": (555, 810, 930, 1000, "#f2e6f0", ["отклонённое действие", "код и причина", "пользователь и время"]),
    }
    for title, (x1, y1, x2, y2, fill, lines) in boxes.items():
        draw_box(d, (x1, y1, x2, y2), title, lines, fill)
    links = [
        ((370, 235), (510, 235)), ((230, 320), (230, 430)), ((410, 540), (510, 580)),
        ((870, 225), (1010, 225)), ((1325, 225), (1435, 225)), ((870, 580), (1010, 580)),
        ((1325, 580), (1435, 580)), ((690, 350), (690, 470)), ((1160, 690), (750, 810)),
    ]
    for a, b in links:
        draw_arrow(d, a, b)
    d.text((900, 760), "Связи обеспечиваются внешними ключами и ограничениями целостности Oracle", font=font(24), fill="#40566f", anchor="mm")
    img.save(ASSETS / "lr1_conceptual_er.png", quality=95)

    # Схема основных таблиц для второй лабораторной работы.
    img = Image.new("RGB", (1800, 1100), "white")
    d = ImageDraw.Draw(img)
    d.text((900, 44), "Основные физические сущности и связи", font=font(38, True), fill="#172b4d", anchor="mm")
    phys = {
        "DURAK_PLAYER": (50, 130, 420, 330, "#f9e2e8", ["PK player_id", "UQ db_username", "display_name, locale"]),
        "DURAK_GAME": (550, 110, 960, 390, "#dcecff", ["PK game_id", "FK creator, daily, trump", "status, phase, seed_hash", "attacker/defender/actor"]),
        "DURAK_GAME_PLAYER": (50, 470, 470, 720, "#e5f5e0", ["PK game_id + seat_no", "FK player_id", "hand_count, status", "finish_place"]),
        "DURAK_GAME_CARD": (550, 490, 960, 750, "#fff1cc", ["PK game_id + card_id", "zone_code, zone_position", "holder_seat_no", "table_pair_no"]),
        "DURAK_ROUND": (1080, 110, 1450, 350, "#eee5ff", ["PK game_id + round_no", "attacker_seat_no", "defender_seat_no", "round_result"]),
        "DURAK_TABLE_PAIR": (1080, 470, 1510, 720, "#e6f3ff", ["PK game_id + round + pair", "attack_game_card_id", "defense_game_card_id", "pair_status"]),
        "DURAK_EVENT": (1080, 820, 1510, 1050, "#fde7d9", ["PK game_id + event_no", "action_type, source", "actor_player_id", "result / state snapshot"]),
        "DURAK_CARD": (50, 840, 430, 1050, "#e4f7f4", ["PK card_id", "UQ card_code", "rank_value, suit_code", "included_in_36"]),
    }
    for title, (x1, y1, x2, y2, fill, lines) in phys.items():
        draw_box(d, (x1, y1, x2, y2), title, lines, fill)
    for a, b in [
        ((420, 230), (550, 230)), ((230, 330), (230, 470)), ((470, 590), (550, 620)),
        ((960, 230), (1080, 230)), ((960, 620), (1080, 600)), ((1295, 350), (1295, 470)),
        ((760, 750), (760, 900)), ((430, 940), (550, 650)), ((1295, 720), (1295, 820)),
    ]:
        draw_arrow(d, a, b)
    img.save(ASSETS / "lr2_physical_er.png", quality=95)

    # Схема взаимодействия пакетов для второй лабораторной работы.
    img = Image.new("RGB", (1800, 1050), "white")
    d = ImageDraw.Draw(img)
    d.text((900, 50), "Архитектура серверной реализации", font=font(40, True), fill="#172b4d", anchor="mm")
    layers = [
        (70, 120, 1730, 250, "#e9f2ff", "SQL-клиенты", ["SQL Developer / SQLcl / SQL*Plus", "команды пользователя и чтение представлений"]),
        (70, 310, 1730, 475, "#e7f7f0", "Публичный интерфейс", ["DURAK_CONSOLE — короткие команды", "DURAK_API — атомарные операции", "V_GAME_CONSOLE, V_LOG_TEXT и защищённые представления"]),
        (70, 535, 1730, 740, "#fff3d6", "Предметная логика", ["DURAK_ENGINE — игровой автомат", "DURAK_RULES — правила боя и лимиты", "DURAK_RANDOM — seed и порядок колоды"]),
        (70, 800, 1730, 980, "#f4e9f6", "Хранение и фоновые процессы", ["13 таблиц Oracle, ограничения, индексы", "DURAK_MAINTENANCE + DBMS_SCHEDULER", "журнал событий, ошибки, Daily и реплей"]),
    ]
    for x1, y1, x2, y2, fill, title, lines in layers:
        d.rounded_rectangle((x1, y1, x2, y2), radius=20, fill=fill, outline="#40566f", width=3)
        d.text((x1 + 34, y1 + 28), title, font=font(27, True), fill="#172b4d")
        for i, line in enumerate(lines):
            d.text((x1 + 610, y1 + 25 + i * 42), "• " + line, font=font(23), fill="#25384c")
    for y1, y2 in [(250, 310), (475, 535), (740, 800)]:
        draw_arrow(d, (900, y1), (900, y2), width=5)
    img.save(ASSETS / "lr2_architecture.png", quality=95)


def empty_reference(reference: Path) -> Document:
    doc = Document(reference)
    body = doc._body._element
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)
    # Очистка связей и вложений исходного документа.
    for rel_id, rel in list(doc.part.rels.items()):
        if rel.reltype in (RT.IMAGE, RT.HYPERLINK):
            doc.part.drop_rel(rel_id)
    for section in doc.sections:
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width = Cm(21)
        section.page_height = Cm(29.7)
        section.left_margin = Cm(3)
        section.right_margin = Cm(1.5)
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.header_distance = Cm(1)
        section.footer_distance = Cm(1)
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(14)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    return doc


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=60, start=80, bottom=60, end=80):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tcMar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement("w:tblHeader")
    tblHeader.set(qn("w:val"), "true")
    trPr.append(tblHeader)


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "start" if edge == "left" else "end" if edge == "right" else edge
        element = OxmlElement(f"w:{tag}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "8")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "000000")
        borders.append(element)


def apply_run(run, size=14, bold=None, italic=None, font_name="Times New Roman"):
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def format_paragraph(p, *, align=WD_ALIGN_PARAGRAPH.JUSTIFY, indent=True, line=1.5, before=0, after=0):
    p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    if indent:
        pf.first_line_indent = Cm(1.25)
    else:
        pf.first_line_indent = None
    pf.widow_control = True


def add_para(doc: Document, text: str = "", *, align=WD_ALIGN_PARAGRAPH.JUSTIFY, indent=True, bold=False, italic=False, size=14, line=1.5, before=0, after=0):
    p = doc.add_paragraph()
    format_paragraph(p, align=align, indent=indent, line=line, before=before, after=after)
    r = p.add_run(text)
    apply_run(r, size=size, bold=bold, italic=italic)
    return p


def add_mixed(doc: Document, parts: Sequence[tuple[str, bool, bool]], *, indent=True, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=14, line=1.5):
    p = doc.add_paragraph()
    format_paragraph(p, align=align, indent=indent, line=line)
    for text, bold, italic in parts:
        r = p.add_run(text)
        apply_run(r, size=size, bold=bold, italic=italic)
    return p


def add_heading(doc: Document, text: str, *, center=False, before=0, after=0):
    p = add_para(doc, text, align=WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT, indent=False, bold=True, before=before, after=after)
    p.paragraph_format.keep_with_next = True
    return p


def add_bullet(doc: Document, text: str, *, level=0, size=14):
    p = doc.add_paragraph()
    format_paragraph(p, indent=False, line=1.5)
    p.paragraph_format.left_indent = Cm(1.25 + 0.75 * level)
    p.paragraph_format.first_line_indent = Cm(-0.5)
    r = p.add_run("• " + text)
    apply_run(r, size=size)
    return p


def add_number(doc: Document, text: str, *, size=14):
    counter = getattr(add_number, "counter", 0) + 1
    add_number.counter = counter
    p = doc.add_paragraph()
    format_paragraph(p, indent=False, line=1.5)
    p.paragraph_format.left_indent = Cm(1.25)
    p.paragraph_format.first_line_indent = Cm(-0.5)
    r = p.add_run(f"{counter}. {text}")
    apply_run(r, size=size)
    return p


def add_page_break(doc: Document):
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)


def add_title_block(doc: Document, number: int):
    add_para(doc, f"Отчёт по Лабораторной работе {number}", align=WD_ALIGN_PARAGRAPH.CENTER, indent=False, bold=True)
    add_para(doc, "по курсу программирование в системах баз данных", align=WD_ALIGN_PARAGRAPH.CENTER, indent=False, bold=True)
    add_para(doc, "", indent=False)
    add_para(doc, "Группа: КА-22-06", align=WD_ALIGN_PARAGRAPH.RIGHT, indent=False)
    add_para(doc, "Выполнил: Михаил", align=WD_ALIGN_PARAGRAPH.RIGHT, indent=False)
    add_para(doc, "", indent=False)


def add_table(doc: Document, headers: Sequence[str], rows: Sequence[Sequence[str]], widths: Sequence[float] | None = None, font_size=10):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_borders(table)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    hdr = table.rows[0]
    repeat_table_header(hdr)
    for i, value in enumerate(headers):
        cell = hdr.cells[i]
        set_cell_shading(cell, "D9EAF7")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell)
        if widths:
            cell.width = Cm(widths[i])
        p = cell.paragraphs[0]
        format_paragraph(p, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False, line=1.0)
        r = p.add_run(value)
        apply_run(r, size=font_size, bold=True)
    for row_values in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row_values):
            cell = cells[i]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if widths:
                cell.width = Cm(widths[i])
            p = cell.paragraphs[0]
            format_paragraph(p, align=WD_ALIGN_PARAGRAPH.LEFT, indent=False, line=1.0)
            r = p.add_run(str(value))
            apply_run(r, size=font_size)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_figure(doc: Document, path: Path, caption: str, width_inches=6.4):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    p.add_run().add_picture(str(path), width=Inches(width_inches))
    c = add_para(doc, caption, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False, size=12, line=1.0)
    c.paragraph_format.keep_with_next = True


def add_code(doc: Document, code: str):
    for line in code.strip("\n").splitlines():
        p = doc.add_paragraph()
        format_paragraph(p, align=WD_ALIGN_PARAGRAPH.LEFT, indent=False, line=1.0)
        p.paragraph_format.left_indent = Cm(0.7)
        r = p.add_run(line)
        apply_run(r, size=9, font_name="Courier New")


def build_lr1() -> Document:
    doc = empty_reference(TEMPLATE_LR1)
    add_title_block(doc, 1)
    add_heading(doc, "Исходное задание")
    add_heading(doc, "Только Oracle: «Дурак»")
    add_heading(doc, "1) Цель")
    add_para(doc, "Создать серверную карточную игру «Дурак», полностью работающую в Oracle Database. Система должна проводить партию от регистрации игроков и детерминированной раздачи до определения проигравшего, сохраняя историю раундов, статистику и результаты. Взаимодействие выполняется через SQL-клиент, представления и текстовые сообщения.")
    add_heading(doc, "2) Ограничения интерфейса")
    add_bullet(doc, "Графический интерфейс и внешние сервисы не используются; источником истины являются таблицы Oracle.")
    add_bullet(doc, "Игровые действия выполняются только через публичные PL/SQL-пакеты; прямой DML в игровые таблицы пользователям не выдаётся.")
    add_bullet(doc, "Скрытые карты видит только их владелец; для остальных участников доступны количество карт, стол, козырь и открытые события.")
    add_heading(doc, "3) Термины и допущения")
    add_bullet(doc, "Колода — 36 карт по умолчанию либо 52 карты; код карты состоит из ранга и масти, например 10H или QC.")
    add_bullet(doc, "Талон — закрытая стопка добора; нижняя открытая карта определяет козырную масть.")
    add_bullet(doc, "Раунд — последовательность атаки, защиты, подбрасывания и завершения отбоем либо взятием.")

    add_page_break(doc)
    add_heading(doc, "4) Уровни требований")
    add_heading(doc, "4.1 Базовый уровень")
    for t in [
        "Партия для двух игроков, подкидной вариант и колода 36 карт.",
        "Корректная раздача по шесть карт, открытый козырь и первый ход по младшему козырю.",
        "Атака, защита, допустимое подбрасывание, завершение раунда и добор до шести карт.",
        "Детерминированная раскладка по seed, история и простое табло результатов.",
    ]:
        add_bullet(doc, t)
    add_heading(doc, "4.2 Хороший уровень")
    for t in [
        "От 2 до 6 игроков, колоды 36/52, подкидной и переводной варианты.",
        "Выбор первого хода по младшему козырю либо псевдослучайно по seed.",
        "Защищённые представления и маскирование чужих рук.",
        "Таймеры действий и простоя, Daily-режим и пошаговый реплей.",
    ]:
        add_bullet(doc, t)
    add_heading(doc, "5) Ключевые бизнес-правила")
    add_bullet(doc, "Карту можно побить старшей картой той же масти; некозырь также бьётся любым козырем, а козырь — только старшим козырем.")
    add_bullet(doc, "Число атакующих пар ограничено max_pairs и количеством карт защитника; для 36 карт значение по умолчанию равно 6, для 52 — 8.")
    add_bullet(doc, "Подбрасывать разрешается только карты достоинств, уже присутствующих на столе.")
    add_bullet(doc, "Перевод возможен до начала защиты и только картой того же достоинства; следующий активный игрок становится защитником.")

    add_page_break(doc)
    add_heading(doc, "6) Пользовательские сценарии")
    scenarios = [
        ("Создание партии", "игрок задаёт колоду, вариант, способ первого хода, seed и таймеры; система возвращает GAME_ID и текст подтверждения."),
        ("Присоединение и старт", "до шести участников занимают места, после чего создатель запускает детерминированную раздачу."),
        ("Игровое действие", "участник атакует, защищается, подбрасывает, переводит, пасует или берёт карты; операция либо атомарно применяется, либо отклоняется с причиной."),
        ("Просмотр состояния", "представления показывают текущую фазу, роли, козырь, талон, свою руку, стол и разрешённые действия."),
        ("Завершение", "после отбоя карты уходят в сброс, после взятия — в руку защитника; затем выполняется добор и начинается новый раунд."),
        ("Daily и реплей", "пользователь получает детерминированную раскладку дня, дневное табло и пошаговую историю партии."),
    ]
    for title, desc in scenarios:
        add_mixed(doc, [(title + ": ", True, False), (desc, False, False)])
    add_heading(doc, "7) Данные и хранение")
    for t in [
        "игроки, Oracle-учётные записи, локаль и часовой пояс;",
        "параметры партий, текущая фаза, роли, seed и причина завершения;",
        "карты партии и их единственная текущая зона: талон, рука, стол или сброс;",
        "раунды, пары атака/защита, журнал событий и перемещений карт;",
        "Daily-конфигурации, результаты дня и ошибки правил.",
    ]:
        add_bullet(doc, t)

    add_page_break(doc)
    add_heading(doc, "8) Отчётность и представления")
    view_rows = [
        ("V_GAME_STATUS", "свод партии по каждому участнику и допустимые действия"),
        ("V_TABLE", "пары атака/защита и состояние текущего стола"),
        ("V_HAND_MINE / V_HAND_PUBLIC", "собственная открытая рука и публичные счётчики чужих рук"),
        ("V_LOG / V_LOG_TEXT", "защищённый и человеко-понятный протокол действий"),
        ("V_REPLAY", "пошаговое воспроизведение событий и перемещений карт"),
        ("V_LEADERBOARD", "победы, поражения и средняя длительность"),
        ("V_DAILY / V_DAILY_RESULTS", "раскладка и итоги дня"),
        ("V_ERRORS", "последние доступные текущему пользователю отказы"),
        ("V_GAME_CONSOLE", "единая текстовая строка состояния для ручной игры"),
    ]
    add_table(doc, ["Представление", "Назначение"], view_rows, [6.2, 9.4], font_size=10)
    add_heading(doc, "9) Нефункциональные требования")
    add_bullet(doc, "Повторная установка должна быть предсказуемой; все сценарии хранятся в UTF-8.")
    add_bullet(doc, "Одинаковые параметры и seed формируют одинаковый порядок колоды.")
    add_bullet(doc, "Одна игровая операция принимается за раз; конфликт блокировки возвращает понятный отказ.")
    add_bullet(doc, "Один Oracle-пользователь не может иметь две активные игровые сессии.")

    add_page_break(doc)
    add_heading(doc, "Перечень основных задач для создания игры «Дурак»", center=True)
    add_para(doc, "Работы разделены на проектирование данных, реализацию игрового автомата, создание защищённого интерфейса и проверку корректности. Такое разделение позволяет независимо тестировать правила карт, порядок раздачи и переходы между фазами.")
    tasks = [
        "Спроектировать нормализованные таблицы и ограничения целостности.",
        "Заполнить справочник из 52 карт и признак принадлежности короткой колоде.",
        "Реализовать стабильное SHA-256-перемешивание и Daily seed.",
        "Реализовать игровой автомат с блокировкой строки партии FOR UPDATE WAIT 5.",
        "Добавить публичный API, человеко-понятную консоль и защищённые представления.",
        "Реализовать таймеры, Daily, реплей и итоговые табло.",
        "Подготовить модульные, интеграционные и конкурентные SQL-тесты.",
    ]
    for t in tasks:
        add_number(doc, t)
    add_heading(doc, "Принятые проектные решения")
    add_para(doc, "Каждая карта партии хранится одной строкой DURAK_GAME_CARD и в любой момент принадлежит ровно одной зоне. Состояние партии хранится в DURAK_GAME, а неизменяемая последовательность событий — в DURAK_EVENT. Логика разделена между пакетами правил, генератора, движка и публичного API.")

    add_page_break(doc)
    add_heading(doc, "Жизненный цикл партии")
    lifecycle = [
        ("LOBBY", "создание партии и подключение игроков; резервируется единственная активная сессия пользователя."),
        ("WAIT_ATTACK", "атакующий кладёт первую карту; при истечении таймера движок выбирает допустимое авто-действие."),
        ("WAIT_DEFENSE", "защитник выбирает карту покрытия, переводит атаку либо принимает решение взять."),
        ("WAIT_THROW", "атакующая сторона по очереди подбрасывает допустимые карты или пасует."),
        ("ROUND_END", "стол уходит в сброс либо в руку защитника, затем игроки добирают карты в установленном порядке."),
        ("FINISHED / EXPIRED", "фиксируется дурак, ничья, отмена или завершение по бездействию; активные сессии освобождаются."),
    ]
    for phase, desc in lifecycle:
        add_mixed(doc, [(phase + " — ", True, False), (desc, False, False)])
    add_heading(doc, "Конкурентность и транзакции")
    add_para(doc, "Публичная операция определяет текущего игрока по SESSION_USER, блокирует строку партии и выполняет всю проверку в одной транзакции. При конфликте ожидание ограничено пятью секундами; вместо двойного хода возвращается сообщение «Партия занята другим действием». Успешные операции и бизнес-отказы фиксируются отдельно.")
    add_heading(doc, "Автоматические действия")
    add_para(doc, "Фоновое задание DBMS_SCHEDULER вызывает DURAK_MAINTENANCE. По таймеру атакующая сторона завершает атаку или пасует, защитник берёт карты, а длительно простаивающая партия переводится в состояние EXPIRED с причиной IDLE_TIMEOUT.")

    add_page_break(doc)
    add_heading(doc, "Проектирование базы данных")
    add_para(doc, "Концептуальная модель связывает участников, партии, карты, раунды и события. Сущности Daily и Error дополняют основной цикл раскладкой дня и журналом отказов. На рисунке 1 показаны главные связи.")
    add_figure(doc, ASSETS / "lr1_conceptual_er.png", "Рисунок 1 — Концептуальная ER-диаграмма игры «Дурак»", 6.35)
    add_para(doc, "Связь PLAYER—GAME_PLAYER является отношением многие-ко-многим между игроками и партиями. GAME_CARD связывает справочник карт с конкретной партией и хранит текущую зону. ROUND и TABLE_PAIR образуют структуру стола, а EVENT обеспечивает аудит и реплей.")
    add_heading(doc, "Целостность")
    add_para(doc, "Первичные и внешние ключи не допускают «осиротевших» записей. CHECK-ограничения фиксируют допустимые статусы, фазы, размеры колоды и варианты правил. Уникальный индекс на DURAK_ACTIVE_SESSION реализует правило одной активной сессии для человека.")

    add_page_break(doc)
    add_heading(doc, "Реализация логики и интерфейса")
    add_heading(doc, "1. Серверное ядро")
    add_para(doc, "DURAK_ENGINE реализует создание партии, раздачу, атаки, защиту, подбрасывание, перевод, взятие, отбой, добор и завершение. DURAK_RULES отвечает за сравнение карт и расчёт лимитов, а DURAK_RANDOM — за нормализацию seed и устойчивый порядок колоды.")
    add_heading(doc, "2. Публичный API")
    add_para(doc, "DURAK_API имеет права владельца схемы и скрывает внутренние таблицы. Для каждой операции возвращается сообщение «УСПЕШНО» либо «ОТКАЗ» с понятной причиной. Пользовательские права ограничены выполнением API и чтением представлений.")
    add_heading(doc, "3. Человеко-понятная консоль")
    add_code(doc, "VARIABLE message VARCHAR2(1000)\nEXEC durak_console.play('ХОД 8H', :message)\nEXEC durak_console.play('БИТО 9H 1', :message)\nEXEC durak_console.play('ПАС', :message)\nPRINT message")
    add_para(doc, "При передаче только кода карты консоль сама определяет действие по текущей фазе. Состояние отображается одним запросом к V_GAME_CONSOLE, а текстовый журнал — запросом к V_LOG_TEXT.")
    add_heading(doc, "4. Защита скрытой информации")
    add_para(doc, "Представления используют SESSION_USER, а не CURRENT_SCHEMA. Поэтому даже после ALTER SESSION SET CURRENT_SCHEMA игрок не получает доступ к чужим картам; приватные события активной партии скрываются до её завершения.")

    add_page_break(doc)
    add_heading(doc, "Проверка выполнения требований")
    checks = [
        ("Справочник карт", "52/52 карты, 36 карт короткой колоды, уникальные коды"),
        ("Правила боя", "масти, козыри, ранги и лимиты стола"),
        ("Seed", "повторяемая раздача и иной порядок для другого seed"),
        ("Видимость", "маскирование руки, лога, реплея и ошибок"),
        ("Игровой автомат", "полный раунд, взятие, отбой, добор и таймеры"),
        ("Хороший уровень", "6 игроков, колода 52, перевод, Daily и реплей"),
        ("Конкурентность", "вторая сессия получает ORA-20702 без двойного действия"),
    ]
    add_table(doc, ["Область", "Результат"], checks, [5.2, 10.4], font_size=10)
    add_para(doc, "Итоговый автоматический прогон завершён без ошибок: 78 проверок из 78 пройдены. Дополнительно подтверждены перевод атаки, режим шести игроков с колодой 52 карты и корректный отказ при конфликте двух одновременных Oracle-сессий.")
    add_heading(doc, "Вывод")
    add_para(doc, "Сформирована полная модель серверной игры «Дурак» в Oracle. Требования декомпозированы на данные, правила, автомат состояний, безопасность и отчётность. Архитектура допускает одновременные независимые партии и сохраняет детерминизм при одинаковых параметрах и действиях.")
    add_heading(doc, "Список использованных источников")
    for source in [
        "1. Техническое задание проекта «Только Oracle: Дурак», предоставленное заказчиком работы.",
        "2. Oracle Database PL/SQL Language Reference, 19c. URL: https://docs.oracle.com/en/database/oracle/oracle-database/19/lnpls/",
        "3. Oracle Database SQL Language Reference, 19c. URL: https://docs.oracle.com/en/database/oracle/oracle-database/19/sqlrf/",
        "4. Исходный код проекта «Игра Дурак только на Oracle»: схема, пакеты, представления и SQL-тесты.",
    ]:
        add_para(doc, source, indent=False, size=12, line=1.0)
    return doc


def build_lr2() -> Document:
    doc = empty_reference(TEMPLATE_LR2)
    add_title_block(doc, 2)
    add_heading(doc, "Исходное задание")
    add_heading(doc, "Только Oracle: «Дурак»")
    add_heading(doc, "1) Цель")
    add_para(doc, "Реализовать в Oracle Database полный серверный цикл карточной игры «Дурак»: создание партии, раздачу, козырь, атаки и защиту, подбрасывание, перевод, взятие, отбой, добор, завершение, историю и табло.")
    add_heading(doc, "2) Реализованные режимы")
    for t in [
        "от 2 до 6 игроков с пропуском вышедших;",
        "колоды 36 и 52 карты; подкидной и переводной варианты;",
        "первый ход по младшему козырю либо детерминированно по seed;",
        "Daily с дневным табло и пошаговый реплей;",
        "таймер хода, автоматические действия и завершение при простое.",
    ]:
        add_bullet(doc, t)
    add_heading(doc, "3) Интерфейс")
    add_para(doc, "Пользователь работает из SQL Developer, SQLcl или SQL*Plus. Операции выполняются через DURAK_API и DURAK_CONSOLE, а состояние читается из представлений. Таблицы владельца схемы недоступны игрокам напрямую.")

    add_page_break(doc)
    add_heading(doc, "4) Основные правила реализации")
    rules = [
        "Защита допустима старшей картой той же масти либо козырем; старший козырь бьёт младший.",
        "В подбрасывании участвуют только достоинства, уже находящиеся на столе.",
        "Лимит пар равен минимуму из max_pairs и исходного количества карт защитника.",
        "Перевод разрешён только до первой защиты и сохраняет ограничения стола.",
        "После отбоя следующий ход начинает защитник; после взятия — следующий активный игрок после защитника.",
        "Добор выполняется сначала атакующей стороной, затем защитником, пока в талоне остаются карты.",
        "Игрок выходит при пустой руке после исчерпания талона; последний игрок с картами объявляется дураком.",
    ]
    for t in rules:
        add_bullet(doc, t)
    add_heading(doc, "5) Транзакционная модель")
    add_para(doc, "Все публичные операции атомарны. DURAK_ENGINE блокирует строку DURAK_GAME конструкцией SELECT … FOR UPDATE WAIT 5. Если другая сессия уже выполняет действие, пользователь получает ORA-20702 с предложением повторить попытку. Версия состояния и последовательный номер события предотвращают неоднозначность журнала.")
    add_heading(doc, "6) Детерминированность")
    add_para(doc, "Порядок колоды строится на STANDARD_HASH с SHA-256. В ключ включаются seed, параметры колоды, вариант правил и идентификатор карты. Поэтому одинаковый набор параметров формирует одинаковую раздачу и талон без зависимости от NLS-настроек Oracle-сессии.")

    add_page_break(doc)
    add_heading(doc, "7) Представления и видимость")
    add_para(doc, "Защищённые представления определяют текущего игрока по SYS_CONTEXT('USERENV','SESSION_USER'). Собственная рука доступна в V_HAND_MINE, а чужие руки представлены только счётчиками в V_HAND_PUBLIC. V_LOG и V_REPLAY скрывают закрытые события активной партии.")
    view_rows = [
        ("V_GAME_STATUS", "параметры, роли, талон, число карт и допустимые действия"),
        ("V_TABLE", "открытые пары атака/защита"),
        ("V_HAND_MINE", "коды карт текущего Oracle-пользователя"),
        ("V_HAND_PUBLIC", "публичные сведения об участниках"),
        ("V_LOG", "фильтрованный протокол"),
        ("V_REPLAY", "пошаговое состояние и перемещения карт"),
        ("V_LEADERBOARD", "свод побед и поражений"),
        ("V_DAILY, V_DAILY_RESULTS", "Daily-конфигурации и места"),
        ("V_ERRORS", "доступные пользователю отказы"),
        ("V_GAME_CONSOLE, V_LOG_TEXT", "текстовый интерфейс игры"),
    ]
    add_table(doc, ["Представление", "Назначение"], view_rows, [6.5, 9.1], font_size=9)
    add_para(doc, "Всего проект создаёт 13 представлений. V_GAME_CONSOLE и V_LOG_TEXT упрощают ручную демонстрацию и не меняют серверные правила.")

    add_page_break(doc)
    add_heading(doc, "Перечень основных объектов, необходимых для игры «Дурак»", center=True)
    add_para(doc, "Физическая схема содержит 13 таблиц, последовательность идентификаторов, индексы, шесть PL/SQL-пакетов, тринадцать представлений и одно фоновое задание DBMS_SCHEDULER.")
    objects = [
        ("Справочники", "DURAK_PLAYER, DURAK_CARD"),
        ("Основной цикл", "DURAK_GAME, DURAK_GAME_PLAYER, DURAK_ACTIVE_SESSION, DURAK_GAME_CARD, DURAK_ROUND, DURAK_TABLE_PAIR"),
        ("Аудит", "DURAK_EVENT, DURAK_CARD_MOVE, DURAK_ERROR"),
        ("Daily", "DURAK_DAILY, DURAK_DAILY_RESULT"),
        ("Пакеты", "RANDOM, RULES, ENGINE, API, CONSOLE, MAINTENANCE"),
        ("Планировщик", "DURAK_MAINTENANCE_JOB"),
    ]
    add_table(doc, ["Группа", "Объекты"], objects, [4.2, 11.4], font_size=10)
    add_heading(doc, "Ответственность слоёв")
    add_para(doc, "Таблицы хранят состояние и журнал, прикладные пакеты изменяют данные, публичный API преобразует ошибки в понятные сообщения, а представления формируют безопасную проекцию. Консоль является тонким адаптером над API и не содержит самостоятельных игровых правил.")
    add_heading(doc, "Индексы")
    add_para(doc, "Индексы созданы для поиска активных партий по дедлайну, состояния игроков, карт по зоне, открытых пар стола, событий по времени и ошибок пользователя. Уникальные ограничения защищают места участников, карты и номера событий.")

    add_page_break(doc)
    add_heading(doc, "I. Сущности и назначение")
    entity_rows = [
        ("DURAK_PLAYER", "игроки", "учётная запись, отображаемое имя, локаль и часовой пояс"),
        ("DURAK_CARD", "справочник карт", "код, ранг, масть, порядок, принадлежность колоде 36"),
        ("DURAK_GAME", "партия", "параметры, seed, фаза, роли, таймеры, результат"),
        ("DURAK_GAME_PLAYER", "участник партии", "место, статус, число карт, место выхода"),
        ("DURAK_ACTIVE_SESSION", "активная сессия", "одна партия на человека"),
        ("DURAK_GAME_CARD", "экземпляр карты", "зона, позиция, владелец, пара стола"),
        ("DURAK_ROUND", "раунд", "атакующий, защитник, лимит, результат"),
        ("DURAK_TABLE_PAIR", "пара стола", "карты атаки и защиты, статус"),
        ("DURAK_EVENT", "протокол", "тип действия, источник, инициатор, снимок состояния"),
        ("DURAK_CARD_MOVE", "движение карты", "зона до/после и закрытость события"),
        ("DURAK_ERROR", "отказ", "код, сообщение, действие и пользователь"),
        ("DURAK_DAILY", "раскладка дня", "дата, правила, seed и хэш конфигурации"),
        ("DURAK_DAILY_RESULT", "результат дня", "место, дурак/не дурак, длительность"),
    ]
    add_table(doc, ["Таблица", "Сущность", "Главные данные"], entity_rows, [4.3, 4.1, 7.2], font_size=8)

    add_page_break(doc)
    add_heading(doc, "Таблица 1 — Ключевые атрибуты основной схемы", center=True)
    attrs = [
        ("DURAK_GAME", "GAME_ID", "NUMBER(19), PK", "идентификатор партии"),
        ("DURAK_GAME", "DECK_SIZE", "NUMBER(2), CHECK", "36 либо 52"),
        ("DURAK_GAME", "GAME_VARIANT", "VARCHAR2(12), CHECK", "PODKIDNOY/PEREVODNOY"),
        ("DURAK_GAME", "SEED_HASH", "RAW(32), NOT NULL", "хэш детерминированной раскладки"),
        ("DURAK_GAME", "GAME_STATUS / PHASE", "VARCHAR2", "статус и фаза автомата"),
        ("DURAK_GAME", "ATTACKER/DEFENDER/ACTOR", "NUMBER(1)", "текущие места участников"),
        ("DURAK_GAME_PLAYER", "GAME_ID + SEAT_NO", "составной PK", "место в партии"),
        ("DURAK_GAME_PLAYER", "PLAYER_ID", "FK", "ссылка на игрока"),
        ("DURAK_GAME_CARD", "GAME_ID + CARD_ID", "составной PK", "экземпляр карты"),
        ("DURAK_GAME_CARD", "ZONE_CODE", "VARCHAR2, CHECK", "талон/рука/стол/сброс"),
        ("DURAK_TABLE_PAIR", "PAIR_NO", "NUMBER", "номер пары атака/защита"),
        ("DURAK_EVENT", "EVENT_NO", "NUMBER", "строгий порядок протокола"),
        ("DURAK_EVENT", "EVENT_SOURCE", "VARCHAR2", "MANUAL/TIMEOUT/SYSTEM/ADMIN"),
        ("DURAK_ERROR", "ERROR_CODE", "NUMBER", "Oracle-код отказа"),
    ]
    add_table(doc, ["Таблица", "Столбец", "Тип/ограничение", "Назначение"], attrs, [3.8, 4.1, 4.0, 4.0], font_size=8)
    add_para(doc, "В Oracle SQL логические значения представлены CHAR(1) со значениями Y/N, а моменты времени — TIMESTAMP WITH TIME ZONE. Это обеспечивает совместимость с Oracle 19c и сохраняет часовой пояс событий.")

    add_page_break(doc)
    add_heading(doc, "II. Процедуры публичного API")
    procedures = [
        ("REGISTER_PLAYER", "регистрация текущего Oracle-пользователя"),
        ("CREATE_GAME / CREATE_DAILY_GAME", "создание обычной или Daily-партии"),
        ("JOIN_GAME / START_GAME", "подключение и запуск раздачи"),
        ("ATTACK", "первая атака раунда"),
        ("DEFEND", "покрытие открытой атакующей карты"),
        ("THROW_IN / PASS_THROW", "подбрасывание или завершение атаки"),
        ("TAKE_CARDS", "взятие карт защитником"),
        ("TRANSFER_ATTACK", "перевод атаки следующему игроку"),
        ("HEARTBEAT / CANCEL_GAME", "активность и явное завершение"),
    ]
    add_table(doc, ["Процедура", "Назначение"], procedures, [7.0, 8.6], font_size=9)
    add_heading(doc, "III. Функции правил и seed")
    functions = [
        ("DURAK_RULES.CAN_BEAT", "проверяет, бьёт ли карта защиты карту атаки"),
        ("DEFAULT_MAX_PAIRS", "возвращает 6 для колоды 36 и 8 для колоды 52"),
        ("EFFECTIVE_ATTACK_LIMIT", "учитывает max_pairs и руку защитника"),
        ("NEXT_ACTIVE_SEAT", "находит следующего не вышедшего игрока"),
        ("RANK_PRESENT_ON_TABLE", "проверяет допустимость достоинства подброса"),
        ("DURAK_RANDOM.DAILY_SEED", "формирует seed по дате и версии правил"),
        ("CARD_SHUFFLE_KEY", "строит SHA-256-ключ карты для сортировки"),
        ("SEEDED_POSITION", "детерминированно выбирает позицию игрока"),
    ]
    add_table(doc, ["Функция", "Назначение"], functions, [7.0, 8.6], font_size=9)

    add_page_break(doc)
    add_heading(doc, "Физическая модель основных таблиц")
    add_figure(doc, ASSETS / "lr2_physical_er.png", "Рисунок 1 — Физическая ER-диаграмма ядра игры", 6.4)
    add_para(doc, "DURAK_GAME является агрегатом состояния партии. DURAK_GAME_PLAYER задаёт участников и места, DURAK_GAME_CARD — расположение карт, DURAK_ROUND и DURAK_TABLE_PAIR — структуру текущего раунда. DURAK_EVENT формирует неизменяемый журнал, а DURAK_CARD_MOVE деталирует перемещение каждой карты.")
    add_para(doc, "Внешние ключи связывают записи внутри одной партии. Составные первичные ключи предотвращают повторение места, карты, номера раунда или номера события. Это позволяет восстанавливать ход партии и проверять согласованность состояния без внешнего приложения.")

    add_page_break(doc)
    add_heading(doc, "Схема программной реализации")
    add_figure(doc, ASSETS / "lr2_architecture.png", "Рисунок 2 — Взаимодействие слоёв и пакетов", 6.4)
    add_para(doc, "Игрок обращается только к публичному слою. DURAK_CONSOLE преобразует русские команды в вызовы DURAK_API. API определяет пользователя, вызывает DURAK_ENGINE и преобразует исключения в текст. Движок использует DURAK_RULES и DURAK_RANDOM; внешняя игровая логика не требуется.")
    add_para(doc, "DURAK_MAINTENANCE запускается планировщиком и обрабатывает истёкшие дедлайны. Представления читают те же таблицы, но фильтруют закрытые данные с учётом SESSION_USER. Тем самым чтение и изменение разделены, а правила не дублируются в клиенте.")

    add_page_break(doc)
    add_heading(doc, "Описание работы основных процедур и функций")
    details = [
        ("1. CREATE_GAME", "создаёт запись DURAK_GAME, нормализует параметры и seed, резервирует активную сессию создателя и добавляет его на первое место. Ошибочные значения колоды, варианта и лимитов отклоняются до изменения состояния."),
        ("2. START_GAME", "блокирует лобби, проверяет число участников, создаёт строки DURAK_GAME_CARD, сортирует их по SHA-256-ключу, раздаёт по шесть карт и выбирает атакующего. Козырь соответствует нижней карте талона."),
        ("3. ATTACK", "проверяет фазу, текущего исполнителя, наличие карты в руке и лимит стола; переносит карту в первую свободную пару и записывает событие ATTACK."),
        ("4. DEFEND", "находит открытую пару, вызывает CAN_BEAT, переносит карту защиты на стол и переводит автомат к следующему допустимому действию."),
        ("5. THROW_IN", "проверяет очередность подбрасывающих, достоинство карты и лимит. После успешного подброса защитник снова получает право защищаться."),
        ("6. PASS_THROW", "фиксирует пас участника. Когда все атакующие завершили подбрасывание и все карты покрыты, запускается отбой и добор."),
    ]
    for title, desc in details:
        add_mixed(doc, [(title + ". ", True, False), (desc, False, False)])
    add_heading(doc, "Пример короткого интерфейса")
    add_code(doc, "SELECT game_text, turn_text, hand_text, table_text,\n       players_text, command_hint, last_event_text\nFROM v_game_console\nWHERE game_id = :game_id;\n\nEXEC durak_console.play('8H', :message)")

    add_page_break(doc)
    add_heading(doc, "Продолжение описания процедур")
    details = [
        ("7. TAKE_CARDS", "переносит все карты стола в руку защитника, фиксирует результат TAKE, выполняет добор атакующей стороны и начинает следующий раунд после защитника."),
        ("8. TRANSFER_ATTACK", "разрешён только до защиты. Защитник добавляет карту того же достоинства, после чего роли циклически сдвигаются к следующему активному игроку."),
        ("9. DRAW_AFTER_ROUND", "добирает карты до целевого размера: сначала атакующие по часовой стрелке, затем защитник. Позиция в талоне уменьшается строго детерминированно."),
        ("10. FINISH_IF_NEEDED", "после исчерпания талона отмечает вышедших игроков. При одном оставшемся игроке записывает дурака, места и причину завершения."),
        ("11. PROCESS_TIMEOUT", "выбирает автоматическое действие по фазе: первую допустимую атаку, пас при подбрасывании или взятие для защитника."),
        ("12. EXPIRE_IDLE_GAME", "переводит простаивающую партию в EXPIRED, записывает IDLE_TIMEOUT и освобождает активные сессии."),
    ]
    for title, desc in details:
        add_mixed(doc, [(title + ". ", True, False), (desc, False, False)])
    add_page_break(doc)
    add_heading(doc, "Безопасность и таймеры")
    add_heading(doc, "Безопасность")
    add_para(doc, "Пакеты объявлены AUTHID DEFINER. Игровые пользователи получают EXECUTE на DURAK_API/DURAK_CONSOLE и SELECT на представления, но не получают прямых прав на таблицы и DURAK_ENGINE. Идентификация всегда выполняется по SESSION_USER.")
    add_heading(doc, "Таймеры")
    add_para(doc, "Задание DURAK_MAINTENANCE_JOB включено в DBMS_SCHEDULER и периодически вызывает RUN_ONCE. Обработчик выбирает только партии с истёкшим action_deadline_at или last_activity_at, затем повторно блокирует строку перед изменением.")
    add_heading(doc, "Обработка ошибок")
    add_para(doc, "Предметные исключения находятся в диапазоне ORA-20xxx. DURAK_API сохраняет отказ через DURAK_ENGINE.LOG_REJECTION и возвращает сообщение без технического стека. V_ERRORS показывает пользователю только собственные ошибки или ошибки доступной партии.")
    add_heading(doc, "Отчуждаемость")
    add_para(doc, "Проект устанавливается сценариями install.sql и reinstall.sql, не требует внешнего сервера и сопровождается демонстрациями для нескольких игроков и Daily. Все объекты создаются в одной схеме-владельце Oracle.")

    add_page_break(doc)
    add_heading(doc, "Тестирование и итоговое назначение системы")
    test_rows = [
        ("Справочник", "8", "карты, масти, ранги и короткая колода"),
        ("Правила", "13", "бой карт, лимиты и входные параметры"),
        ("Seed", "6", "повторяемость порядка и Daily"),
        ("Представления", "13", "маскирование, консоль и реплей"),
        ("Объекты", "4", "валидность, состав и планировщик"),
        ("Игровой автомат", "20", "раунды, взятие, добор, таймеры и сессии"),
        ("Daily", "5", "seed, результаты и табло дня"),
        ("Режимы", "9", "6 игроков, колода 52, перевод и подброс"),
        ("Итого", "78", "78 из 78 проверок пройдены"),
    ]
    add_table(doc, ["Раздел", "Проверок", "Покрытие"], test_rows, [4.4, 2.8, 8.4], font_size=9)
    add_para(doc, "Проверки подтвердили понятность V_GAME_CONSOLE и V_LOG_TEXT, перевод атаки, шестиместную партию с колодой 52 карты и корректный отказ при конкурентном действии. После тестов все пакеты и представления имеют статус VALID, а фоновое задание включено.")
    add_heading(doc, "Вывод")
    add_para(doc, "Пакеты проекта образуют завершённый серверный игровой движок. Схема сохраняет текущее состояние и полный протокол, API обеспечивает атомарность и безопасность, консоль делает управление понятным человеку, а тесты подтверждают выполнение требований минимального и хорошего уровней.")
    add_heading(doc, "Список использованных источников")
    for source in [
        "1. Техническое задание проекта «Только Oracle: Дурак», предоставленное заказчиком работы.",
        "2. Oracle Database PL/SQL Language Reference, 19c. URL: https://docs.oracle.com/en/database/oracle/oracle-database/19/lnpls/",
        "3. Oracle Database SQL Language Reference, 19c. URL: https://docs.oracle.com/en/database/oracle/oracle-database/19/sqlrf/",
        "4. Oracle Database PL/SQL Packages and Types Reference: DBMS_SCHEDULER. URL: https://docs.oracle.com/en/database/oracle/oracle-database/19/arpls/DBMS_SCHEDULER.html",
        "5. Исходный код и тестовые сценарии проекта «Игра Дурак только на Oracle».",
    ]:
        add_para(doc, source, indent=False, size=12, line=1.0)
    return doc


def main() -> None:
    create_diagrams()
    lr1 = build_lr1()
    lr1.core_properties.title = "Лабораторная работа 1 — Дурак на Oracle"
    lr1.core_properties.author = "Михаил"
    lr1.core_properties.last_modified_by = "Михаил"
    lr1.save(OUT_LR1)

    lr2 = build_lr2()
    lr2.core_properties.title = "Лабораторная работа 2 — Дурак на Oracle"
    lr2.core_properties.author = "Михаил"
    lr2.core_properties.last_modified_by = "Михаил"
    lr2.save(OUT_LR2)
    print(OUT_LR1)
    print(OUT_LR2)


if __name__ == "__main__":
    main()
